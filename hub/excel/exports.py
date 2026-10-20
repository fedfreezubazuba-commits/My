"""
Дополнительные выгрузки (раздел 9.5 ТЗ) — не основной путь доставки в модели.

  «Открыть в Excel»  — разовый файл в Выгрузки/ с листами «Данные» (строка на значение, все поля
                       отдельными столбцами, свежие даты сверху) и «Паспорт»; файл сразу открывается.
  Файлы по показателям (настройка, по умолчанию выключена) — постоянный
                       Выгрузки/Показатели/<показатель>.xlsx, обновляется после загрузки.
"""
import datetime as dt
import os
import subprocess
import sys
from pathlib import Path

from ..core import periods as P
from ..core.errors import FileBusy
from ..core.store import series_dims
from . import packet


def dims_table(store, codes):
    """({код: {поле: значение}}, [поля по порядку]) — разрезы рядов как отдельные столбцы."""
    per_code, fields = {}, []
    for c in codes:
        d = dict(series_dims(store.series(c)))
        per_code[c] = d
        for k in d:
            if k not in fields:
                fields.append(k)
    return per_code, fields


def _book(store, codes, passport):
    """Лист «Данные»: строка на значение со всеми полями (свежие даты сверху, фильтр по любому столбцу)
    и лист «Паспорт»."""
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    bold = Font(bold=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Данные"
    meta = {c: store.series(c) for c in codes}
    dims, fields = dims_table(store, codes)
    ws.append(["Дата", "Период"] + fields + ["Периодичность", "Единицы", "Код", "Обновлено", "Значение"])
    rows = []
    for c in codes:
        cut = [dims[c].get(f, "") for f in fields]
        per = P.PERIODICITY_RU.get(meta[c]["periodicity"], meta[c]["periodicity"] or "")
        for o in store.current_rows(c):
            got = (o["fetched_at"] or "")[:10]
            rows.append([dt.date.fromisoformat(o["period_end"]), o["period_label"]] + cut +
                        [per, meta[c]["unit"], c, dt.date.fromisoformat(got) if got else None, o["value"]])
    rows.sort(key=lambda r: r[0], reverse=True)          # последние значения — вверху
    for r in rows:
        ws.append(r)
    for c in ws[1]:
        c.font = bold
    n = len(fields)
    for row in ws.iter_rows(min_row=2):
        row[0].number_format = "DD.MM.YYYY"
        row[n + 5].number_format = "DD.MM.YYYY"
    ws.column_dimensions["A"].width = 12
    for i in range(3, n + 3):
        ws.column_dimensions[get_column_letter(i)].width = 30
    ws.column_dimensions[get_column_letter(n + 5)].width = 26
    ws.auto_filter.ref = ws.dimensions
    ws.freeze_panes = "A2"

    ws = wb.create_sheet("Паспорт")
    for k, v in passport:
        ws.append([k, v])
        ws.cell(ws.max_row, 1).font = bold
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 100
    return wb


def _passport(cfg, store, job, codes):
    from ..sources import get
    src = None
    try:
        src = get(job["source"])
    except KeyError:
        pass
    st = store.state(job["id"])
    rows = [("Показатель", job["name"]), ("Источник", src.title if src else job["source"]),
            ("Ссылка", src.open_url(job.get("params", {})) if src else ""),
            ("Параметры", str(job.get("params", {}))), ("Последняя загрузка", st.get("last_success") or ""),
            ("Выгружено", dt.datetime.now().strftime("%Y-%m-%d %H:%M")),
            ("Дата", "Конец отчётного периода (для снимков — дата наблюдения)")]
    for c in codes:
        s = store.series(c)
        rows.append((c, f"{s['name']} [{s['unit']}]; правило даты: {s['date_rule'] or '—'}"))
    return rows


def _safe_save(wb, path: Path):
    path.parent.mkdir(parents=True, exist_ok=True)
    try:
        wb.save(path)
    except PermissionError:
        raise FileBusy(path)
    return path


def export_job(cfg, store, paths, job_id, codes=None) -> Path:
    job = cfg.job(job_id) or {"id": job_id, "name": job_id, "source": "", "params": {}}
    codes = codes or [r["code"] for r in store.series_of_job(job_id)]
    wb = _book(store, codes, _passport(cfg, store, job, codes))
    name = packet.safe_name(job["name"])[:60]
    return _safe_save(wb, paths.exports / f"{name} {dt.datetime.now():%Y-%m-%d_%H%M%S}.xlsx")


def export_codes(cfg, store, paths, codes, title) -> Path:
    job_id = store.series(codes[0])["job_id"]
    job = cfg.job(job_id) or {"id": job_id, "name": title, "source": "", "params": {}}
    wb = _book(store, codes, _passport(cfg, store, job, codes))
    return _safe_save(wb, paths.exports / f"{packet.safe_name(title)[:60]} {dt.datetime.now():%Y-%m-%d_%H%M%S}.xlsx")


def update_job_files(cfg, store, paths, jobs) -> dict:
    out = {}
    for job in jobs:
        codes = [r["code"] for r in store.series_of_job(job["id"])]
        if not codes:
            continue
        wb = _book(store, codes, _passport(cfg, store, job, codes))
        try:
            out[job["id"]] = str(_safe_save(wb, paths.job_files / f"{packet.safe_name(job['name'])}.xlsx"))
        except FileBusy as e:
            out[job["id"]] = f"busy: {e}"
    return out


def open_file(path) -> None:
    path = str(path)
    if sys.platform.startswith("win"):
        os.startfile(path)  # noqa: S606 — открыть файл программой по умолчанию
    elif sys.platform == "darwin":
        subprocess.Popen(["open", path])
    else:
        subprocess.Popen(["xdg-open", path], stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
