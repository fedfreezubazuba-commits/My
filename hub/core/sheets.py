"""
«Листы файла в модель» — лист Excel (или CSV) переносится в модель как есть, на постоянное имя листа, без разбора
в базу. Два режима:

  replace — файл каждый раз содержит всю историю: в модели всегда последняя версия листа;
  append  — файл содержит недавние периоды (скользящие два квартала): каждое обновление дописывается под
            предыдущими, первые столбцы — «Обновлено» и «Файл от».

Для формул рядом может быть лист «<имя> (ячейки)»: по одной строке на каждое число листа —
Обновлено | Файл от | Лист в файле | Строка | Подпись | Путь | Столбец | Буква | Заголовок | Значение | Таблица.
«Путь» — подпись строки вместе с вышестоящими по отступам («Счет текущих операций / Товары / Экспорт»),
«Заголовок» — шапка столбца, склеенная по этажам и объединённым ячейкам («Справочно: — январь - август 2025 г.»).
Тогда последнее известное значение одной постоянной формулой: СУММЕСЛИМН по Путь + Заголовок + Обновлено =
МАКСЕСЛИ(…) — работает, даже когда столбцы в новых версиях файла сдвигаются.

Хранилище (не база): <рабочая папка>/Листы/<показатель>/<имя>.csv, «<имя> — ячейки.csv», state.json (хэш
содержимого — чтобы не дописывать ту же версию повторно). В «Для моделей» хаб копирует их рядом с посылкой
модели: «<модель> — лист <имя>.csv» и «<модель> — лист <имя> — ячейки.csv» (их читают запросы Power Query).
"""
import csv
import datetime as dt
import hashlib
import io
import json
import re
from pathlib import Path

from . import periods as P
from .errors import StructureError

SHEET_COLS_APPEND = ["Обновлено", "Файл от", "Строка"]
CELL_COLS = ["Обновлено", "Файл от", "Лист в файле", "Строка", "Подпись", "Путь", "Столбец", "Буква", "Заголовок",
             "Значение", "Таблица"]
MAX_NAME = 21               # имя листа в модели; + « (ячейки)» — в пределах 31 знака Excel
BAD_NAME = re.compile(r"[\[\]:*?/\\]")
MODES = {"replace": "Заменять лист: файл каждый раз содержит всю историю — в модели всегда последняя версия",
         "append": "Копить обновления: файл содержит недавние периоды — каждое обновление дописывается под "
                   "предыдущими, с датой обновления"}


# --- чтение листов --------------------------------------------------------------------------
class Sheet:
    """Лист как есть: grid — строки значений (число, текст, дата ISO, None); merges — объединённые ячейки
    (r1, c1, r2, c2) с нуля; indents — {(r, c): отступ} у текстовых ячеек."""

    def __init__(self, title, index, grid, merges=(), indents=None):
        self.title, self.index, self.grid = title, index, grid
        self.merges, self.indents = list(merges), dict(indents or {})


def _value(v):
    if v is None:
        return None
    if isinstance(v, bool):
        return int(v)
    if isinstance(v, (int, float)):
        return int(v) if isinstance(v, float) and v.is_integer() and abs(v) < 1e15 else v
    if isinstance(v, dt.datetime):
        return v.date().isoformat() if v.time() == dt.time() else v.isoformat(sep=" ", timespec="minutes")
    if isinstance(v, dt.date):
        return v.isoformat()
    s = str(v).replace("\r\n", "\n").strip()
    return s or None


def _trim(grid):
    grid = [[_value(c) for c in r] for r in grid]
    while grid and not any(c is not None for c in grid[-1]):
        grid.pop()
    width = max((max((j + 1 for j, c in enumerate(r) if c is not None), default=0) for r in grid), default=0)
    return [list(r[:width]) + [None] * (width - len(r[:width])) for r in grid]


def read_sheets(data: bytes, fmt: str) -> list[Sheet]:
    """Все листы файла как есть (xlsx/xlsm, xls, csv)."""
    out = []
    if fmt == "xlsx":
        from openpyxl import load_workbook
        big = len(data) > 15_000_000          # большой файл — без объединений и отступов (быстрее, меньше памяти)
        wb = load_workbook(io.BytesIO(data), read_only=big, data_only=True)
        for i, ws in enumerate(wb.worksheets):
            if big:
                out.append(Sheet(ws.title, i, _trim(ws.iter_rows(values_only=True))))
                continue
            grid = _trim(ws.iter_rows(values_only=True))
            indents = {}
            for r, row in enumerate(ws.iter_rows(max_row=len(grid), max_col=3)):
                for cell in row:
                    if isinstance(cell.value, str) and cell.alignment is not None and cell.alignment.indent:
                        indents[(r, cell.column - 1)] = int(cell.alignment.indent)
            merges = [(m.min_row - 1, m.min_col - 1, m.max_row - 1, m.max_col - 1) for m in ws.merged_cells.ranges]
            out.append(Sheet(ws.title, i, grid, merges, indents))
    elif fmt == "xls":
        import xlrd
        try:
            wb = xlrd.open_workbook(file_contents=data, formatting_info=True)
        except Exception:  # noqa: BLE001 — без оформления (некоторые файлы его не дают)
            wb = xlrd.open_workbook(file_contents=data)
        for i, sh in enumerate(wb.sheets()):
            rows, indents = [], {}
            for r in range(sh.nrows):
                row = []
                for c in range(sh.ncols):
                    cell = sh.cell(r, c)
                    if cell.ctype == xlrd.XL_CELL_DATE:
                        row.append(xlrd.xldate.xldate_as_datetime(cell.value, wb.datemode))
                    elif cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                        row.append(None)
                    else:
                        row.append(cell.value)
                    if c < 3 and cell.ctype == xlrd.XL_CELL_TEXT and getattr(wb, "xf_list", None):
                        try:
                            ind = wb.xf_list[cell.xf_index].alignment.indent_level
                            if ind:
                                indents[(r, c)] = int(ind)
                        except (IndexError, AttributeError):
                            pass
                rows.append(row)
            merges = [(r1, c1, r2 - 1, c2 - 1) for r1, r2, c1, c2 in getattr(sh, "merged_cells", [])]
            out.append(Sheet(sh.name, i, _trim(rows), merges, indents))
    elif fmt == "csv":
        text = None
        for enc in ("utf-8-sig", "cp1251"):
            try:
                text = data.decode(enc)
                break
            except UnicodeDecodeError:
                continue
        if text is None:
            raise StructureError("Файл не похож на таблицу (не удалось прочитать текст)")
        delim = max([";", ",", "\t", "|"], key=text[:5000].count)
        rows = [[_csv_value(c) for c in r] for r in csv.reader(io.StringIO(text), delimiter=delim)]
        out.append(Sheet("CSV", 0, _trim(rows)))
    else:
        raise StructureError("Перенести листы в модель можно из Excel (xlsx, xls) или CSV")
    out = [s for s in out if s.grid]
    if not out:
        raise StructureError("В файле нет листов с данными")
    return out


NUM_TEXT = re.compile(r"^-?\d[\d\s ]*([.,]\d+)?$")


def _csv_value(s):
    s = (s or "").strip()
    if NUM_TEXT.match(s):
        v = P.parse_number(s)
        if v is not None:
            return v
    return s or None


def bare(title) -> str:
    """Имя листа без цифр и знаков — «Оценка ПБ (15.10.2026)» и «Оценка ПБ (12.11.2026)» совпадают."""
    return re.sub(r"[\W\d_]+", "", str(title or "").lower())


def match_sheet(sheets: list[Sheet], spec: dict) -> Sheet | None:
    """Лист из нового файла, соответствующий выбранному: то же имя → то же имя без дат и цифр → тот же номер
    (если листов столько же) → единственный лист."""
    title = spec.get("title") or ""
    for s in sheets:
        if s.title == title:
            return s
    same = [s for s in sheets if bare(s.title) == bare(title) and bare(title)]
    if len(same) == 1:
        return same[0]
    if same:
        return min(same, key=lambda s: abs(s.index - int(spec.get("index") or 0)))
    idx = spec.get("index")
    if idx is not None and int(idx) < len(sheets) and spec.get("count") == len(sheets):
        return sheets[int(idx)]
    return sheets[0] if len(sheets) == 1 else None


def default_name(title, taken=()) -> str:
    """Постоянное имя листа в модели по имени листа файла: без дат, недопустимых знаков, ≤ 21 знака."""
    s = re.sub(r"\(\s*[\d.\-/ ]+\s*\)|\d{1,2}[./]\d{1,2}[./]\d{2,4}", "", str(title or ""))
    s = BAD_NAME.sub(" ", s)
    s = re.sub(r"\s+", "_", s.strip(" _-'"))[:MAX_NAME].strip("_") or "Лист"
    name, k = s, 2
    while name in taken:
        suffix = f"_{k}"
        name, k = s[:MAX_NAME - len(suffix)] + suffix, k + 1
    return name


def valid_name(name) -> str | None:
    """Почему имя листа не подходит (None — подходит)."""
    if not name or not name.strip():
        return "Укажите имя листа"
    if len(name) > MAX_NAME:
        return f"Имя листа — не длиннее {MAX_NAME} знаков"
    if BAD_NAME.search(name) or name.startswith("'") or name.endswith("'"):
        return "В имени листа нельзя использовать [ ] : * ? / \\ и апостроф по краям"
    if name.lower().startswith("хаб"):
        return "Имена «Хаб…» заняты служебными листами хаба"
    return None


# --- ячейки для формул ----------------------------------------------------------------------
def col_letter(n: int) -> str:
    s = ""
    while n:
        n, r = divmod(n - 1, 26)
        s = chr(65 + r) + s
    return s


def _is_num(v) -> bool:
    return isinstance(v, (int, float)) and not isinstance(v, bool)


def _is_year_row(row) -> bool:
    """Строка шапки из годов (2024 | 2025 | 2026) — числа, но это заголовки."""
    nums = [v for v in row if _is_num(v)]
    return len(nums) >= 2 and all(isinstance(v, int) and 1900 <= v <= 2100 for v in nums) and \
        not any(isinstance(v, str) and P.parse_number(v) is None for v in row[1:] if v)


def _merged_value(sheet: Sheet):
    """{(r, c): значение верхней левой ячейки объединения} для всех ячеек объединений."""
    out = {}
    for r1, c1, r2, c2 in sheet.merges:
        v = sheet.grid[r1][c1] if r1 < len(sheet.grid) and c1 < len(sheet.grid[r1]) else None
        for r in range(r1, r2 + 1):
            for c in range(c1, c2 + 1):
                out[(r, c)] = v
    return out


def _text(v) -> str:
    return " ".join(str(v).split()) if v is not None else ""


def cells(sheet: Sheet) -> list[dict]:
    """Каждое число листа отдельной записью: строка, подпись, путь по отступам, столбец, заголовок, таблица."""
    grid = sheet.grid
    merged = _merged_value(sheet)
    has_indent = bool(sheet.indents)
    title = ""
    header, pending = {}, []
    stack = []                       # [(отступ, подпись)] — вышестоящие строки
    seen = {}
    out = []
    for r, row in enumerate(grid):
        nums = [c for c, v in enumerate(row) if _is_num(v)]
        if not nums or _is_year_row(row):
            filled = [c for c, v in enumerate(row) if v is not None or merged.get((r, c)) is not None]
            if len(filled) == 1 and filled[0] == 0 and not out and not pending and not title:
                title = _text(row[0])                 # строка-название над таблицей
                continue
            if filled:
                pending.append(r)
            continue
        if pending:                                   # шапка из строк над данными
            def filled(p):
                return sum(1 for c in range(len(grid[p])) if grid[p][c] is not None or merged.get((p, c)) is not None)
            rows = [p for p in pending if filled(p) >= 2]
            for p in pending:                         # одна надпись в строке («(млрд долларов США)») — к названию
                if filled(p) == 1 and not header:
                    note = next(_text(v) for v in grid[p] if v is not None) if any(v is not None for v in grid[p]) else ""
                    if note and note not in title:
                        title = f"{title} {note}".strip()
            if rows:
                header = {}
                width = max(len(grid[p]) for p in rows)
                for c in range(width):
                    parts = []
                    for p in rows:
                        v = grid[p][c] if c < len(grid[p]) else None
                        v = v if v is not None else merged.get((p, c))
                        t = _text(v)
                        if t and (not parts or parts[-1] != t):
                            parts.append(t)
                    header[c] = " — ".join(parts)
            pending = []
        first_num = nums[0]
        label_c = next((c for c in range(first_num) if isinstance(row[c], str) and row[c].strip()), None)
        label = _text(row[label_c]) if label_c is not None else ""
        if label:
            if has_indent:
                ind = sheet.indents.get((r, label_c), 0)
            else:
                ind = len(str(row[label_c])) - len(str(row[label_c]).lstrip(" "))
            while stack and stack[-1][0] >= ind:
                stack.pop()
            path = " / ".join([t for _i, t in stack] + [label])
            stack.append((ind, label))
            if not has_indent and not any(i for i, _t in stack[:-1]):
                n = seen.get(label, 0) + 1            # отступов нет — одинаковые подписи нумеруются
                seen[label] = n
                path = label if n == 1 else f"{label} ({n})"
        else:
            path = ""
        for c in nums:
            out.append({"Строка": r + 1, "Подпись": label, "Путь": path, "Столбец": c + 1, "Буква": col_letter(c + 1),
                        "Заголовок": header.get(c, ""), "Значение": row[c], "Таблица": title})
    return out


DATE_TEXT = re.compile(r"(обновл|по состоянию|дата|as of|updated)", re.I)


def file_date(sheet: Sheet) -> str:
    """Дата версии файла: из имени листа («Оценка ПБ (15.10.2026)») или из текста «Дата последнего
    обновления: 15 октября 2026 года». ISO или ""."""
    m = re.search(r"\d{1,2}[./]\d{1,2}[./]\d{4}", sheet.title)
    if m:
        d = P.parse_date(m.group(0))
        if d:
            return d.isoformat()
    for row in sheet.grid:
        for v in row:
            if isinstance(v, str) and DATE_TEXT.search(v):
                d = P.parse_date(v)
                if d:
                    return d.isoformat()
    return ""


# --- хранилище «Листы/» ---------------------------------------------------------------------
def folder(paths, job_id) -> Path:
    return paths.root / "Листы" / job_id


def sheet_csv(paths, job_id, name) -> Path:
    return folder(paths, job_id) / f"{name}.csv"


def cells_csv(paths, job_id, name) -> Path:
    return folder(paths, job_id) / f"{name} — ячейки.csv"


def _num(v) -> str:
    from ..excel.packet import num
    return num(v) if _is_num(v) else ("" if v is None else str(v))


def write_csv(path: Path, header, rows):
    buf = io.StringIO()
    w = csv.writer(buf, delimiter=";", lineterminator="\r\n")
    w.writerow(header)
    for r in rows:
        w.writerow([_num(v) for v in r])
    from ..excel.packet import write_if_changed
    write_if_changed(path, ("﻿" + buf.getvalue()).encode("utf-8"))


def read_csv(path: Path):
    """(заголовок, строки) или (None, [])."""
    if not path.exists():
        return None, []
    rows = list(csv.reader(io.StringIO(path.read_text(encoding="utf-8-sig")), delimiter=";"))
    return (rows[0], rows[1:]) if rows else (None, [])


def _state(paths, job_id) -> dict:
    p = folder(paths, job_id) / "state.json"
    try:
        return json.loads(p.read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return {}


def _letters(width):
    return [col_letter(c + 1) for c in range(width)]


def save(paths, job, got: list[tuple[dict, Sheet]], fetched_at: str) -> list[dict] | None:
    """Записать листы показателя. got — [(описание листа из params.sheets, лист файла)].
    None — содержимое не изменилось с прошлого раза (ничего не пишется). → сводка по листам."""
    params = job.get("params") or {}
    mode = params.get("mode") or "replace"
    digest = hashlib.sha1(json.dumps([(spec["name"], s.grid) for spec, s in got], ensure_ascii=False,
                                     default=str).encode("utf-8")).hexdigest()
    st = _state(paths, job["id"])
    if st.get("hash") == digest and all(sheet_csv(paths, job["id"], spec["name"]).exists() for spec, _s in got):
        return None
    when = fetched_at.replace("T", " ")[:19]
    summary = []
    for spec, sh in got:
        name = spec["name"]
        width = max((len(r) for r in sh.grid), default=0)
        fdate = file_date(sh)
        grid_rows = [[i + 1] + list(r) + [None] * (width - len(r)) for i, r in enumerate(sh.grid)]
        cell_rows = [[when, fdate, sh.title, c["Строка"], c["Подпись"], c["Путь"], c["Столбец"], c["Буква"],
                      c["Заголовок"], c["Значение"], c["Таблица"]] for c in cells(sh)]
        if mode == "append":
            head, old = read_csv(sheet_csv(paths, job["id"], name))
            old_width = (len(head) - len(SHEET_COLS_APPEND)) if head else 0
            w = max(width, old_width)
            rows = [r + [""] * (len(SHEET_COLS_APPEND) + w - len(r)) for r in old]
            rows += [[when, fdate] + r + [None] * (w + 1 - len(r)) for r in grid_rows]
            write_csv(sheet_csv(paths, job["id"], name), SHEET_COLS_APPEND + _letters(w), rows)
            if spec.get("cells", True):
                _h, old_cells = read_csv(cells_csv(paths, job["id"], name))
                write_csv(cells_csv(paths, job["id"], name), CELL_COLS, old_cells + cell_rows)
        else:
            write_csv(sheet_csv(paths, job["id"], name), ["Строка"] + _letters(width), grid_rows)
            if spec.get("cells", True):
                write_csv(cells_csv(paths, job["id"], name), CELL_COLS, cell_rows)
        summary.append({"name": name, "title": sh.title, "rows": len(sh.grid), "cells": len(cell_rows),
                        "file_date": fdate, "mode": mode})
    p = folder(paths, job["id"]) / "state.json"
    p.parent.mkdir(parents=True, exist_ok=True)
    p.write_text(json.dumps({"hash": digest, "updated": fetched_at, "sheets": summary}, ensure_ascii=False),
                 encoding="utf-8")
    return summary


def specs(job) -> list[dict]:
    return list((job.get("params") or {}).get("sheets") or [])


def model_sheets(cfg, model) -> list[dict]:
    """Листы, которые получает модель: [{name, cells, job}] (повторы имён — один раз, первый)."""
    out, seen = [], set()
    for job_id in model.get("sheets") or []:
        job = cfg.job(job_id)
        if not job:
            continue
        for spec in specs(job):
            if spec["name"] in seen:
                continue
            seen.add(spec["name"])
            out.append({"name": spec["name"], "cells": spec.get("cells", True), "job": job_id,
                        "mode": (job.get("params") or {}).get("mode") or "replace"})
    return out


def table_safe(name) -> str:
    s = re.sub(r"[^\w]", "_", str(name))
    return s if s and not s[0].isdigit() else "_" + s


def example_formula(name, cell: dict | None, locale="ru") -> str:
    """Формула «последнее известное значение» для ячейки (строка по пути, столбец по заголовку)."""
    t = f"Ячейки_{table_safe(name)}"
    path = (cell or {}).get("Путь") or "Подпись строки"
    head = (cell or {}).get("Заголовок") or "Заголовок столбца"
    path_q, head_q = path.replace('"', '""'), head.replace('"', '""')
    if locale == "en":
        cond = f'{t}[Путь];"{path_q}";{t}[Заголовок];"{head_q}"'.replace(";", ",")
        return f"=SUMIFS({t}[Значение],{cond},{t}[Обновлено],MAXIFS({t}[Обновлено],{cond}))"
    cond = f'{t}[Путь];"{path_q}";{t}[Заголовок];"{head_q}"'
    return f"=СУММЕСЛИМН({t}[Значение];{cond};{t}[Обновлено];МАКСЕСЛИ({t}[Обновлено];{cond}))"
