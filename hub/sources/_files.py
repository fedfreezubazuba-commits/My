"""
Чтение таблиц из файлов: xlsx, xls, csv, pdf. Каждый лист / таблица PDF становится Table
(как таблица на странице), дальше работает общий разбор из _tables.
"""
import csv
import datetime as dt
import io
import re

from ..core.errors import StructureError
from ._tables import Table, _is_num


def detect_format(data: bytes, name: str = "") -> str:
    n = (name or "").lower()
    if data[:4] == b"%PDF" or n.endswith(".pdf"):
        return "pdf"
    if data[:2] == b"PK":
        return "xlsx" if _is_xlsx(data) else "zip"       # xlsx — тоже zip-архив, но со своей начинкой
    if n.endswith((".xlsx", ".xlsm")):
        return "xlsx"
    if n.endswith(".zip"):
        return "zip"
    if data[:8] == b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" or n.endswith(".xls"):
        return "xls"
    return "csv"


def _is_xlsx(data: bytes) -> bool:
    import zipfile
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            return "[Content_Types].xml" in z.namelist()
    except zipfile.BadZipFile:
        return False


INNER = (".csv", ".txt", ".xlsx", ".xlsm", ".xls", ".pdf")
MAX_UNZIPPED = 1_500_000_000      # больше — скорее всего не таблица, а что-то другое


def zip_members(data: bytes) -> list[tuple[str, bytes]]:
    """[(имя, байты)] файлов с таблицами внутри архива (папки и служебные файлы пропускаются)."""
    import zipfile
    out = []
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        infos = [i for i in z.infolist() if not i.is_dir() and i.filename.lower().endswith(INNER)
                 and not i.filename.split("/")[-1].startswith((".", "__MACOSX", "~$"))
                 and "__MACOSX/" not in i.filename]
        if sum(i.file_size for i in infos) > MAX_UNZIPPED:
            raise StructureError("Архив слишком большой: больше 1,5 ГБ после распаковки")
        for i in infos:
            out.append((i.filename.split("/")[-1], z.read(i)))
    if not out:
        raise StructureError("В архиве нет файлов с таблицами (CSV, Excel, PDF)")
    return out


def _cell(v) -> str:
    if v is None:
        return ""
    if isinstance(v, dt.datetime):
        return v.date().isoformat() if v.time() == dt.time() else v.isoformat(sep=" ", timespec="minutes")
    if isinstance(v, dt.date):
        return v.isoformat()
    if isinstance(v, float):
        return repr(v)[:-2] if v.is_integer() else repr(v)
    return " ".join(str(v).split())


def _to_table(rows, title, index, kind="file", header_row=None):
    rows = [[_cell(c) for c in r] for r in rows]
    rows = [r for r in rows if any(c for c in r)]
    if not rows:
        return None
    width = max(len(r) for r in rows)
    rows = [r + [""] * (width - len(r)) for r in rows]
    if header_row:
        h = int(header_row) - 1
        header, data = rows[h], rows[h + 1:]
    else:
        header, data = [], rows
        for i, r in enumerate(rows[:15]):
            nums = sum(_is_num(c) for c in r)
            if nums == 0 and sum(bool(c) for c in r) >= max(2, width // 2):
                header, data = r, rows[i + 1:]
            elif header and nums:
                break
    if not data:
        return None
    return Table(kind, data, title, index, header)


def read_tables(data: bytes, fmt: str, spec=None):
    spec = spec or {}
    fmt = fmt or "csv"
    tables = []
    if fmt == "zip":
        # архив: таблицы всех файлов внутри; название — имя файла (и листа для Excel)
        members = zip_members(data)
        for name, blob in members:
            for t in read_tables(blob, detect_format(blob, name), spec):
                t.title = name if t.title in ("CSV", "") else f"{name} — {t.title}"
                t.index = len(tables)
                t.archive = True
                tables.append(t)
        return tables
    if fmt == "xlsx":
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        for i, ws in enumerate(wb.worksheets):
            t = _to_table(ws.iter_rows(values_only=True), ws.title, i, header_row=spec.get("header_row"))
            if t:
                tables.append(t)
    elif fmt == "xls":
        import xlrd
        wb = xlrd.open_workbook(file_contents=data)
        for i, sh in enumerate(wb.sheets()):
            rows = []
            for r in range(sh.nrows):
                row = []
                for c in range(sh.ncols):
                    cell = sh.cell(r, c)
                    if cell.ctype == xlrd.XL_CELL_DATE:
                        row.append(xlrd.xldate.xldate_as_datetime(cell.value, wb.datemode))
                    else:
                        row.append(cell.value)
                rows.append(row)
            t = _to_table(rows, sh.name, i, header_row=spec.get("header_row"))
            if t:
                tables.append(t)
    elif fmt == "pdf":
        import pdfplumber
        with pdfplumber.open(io.BytesIO(data)) as pdf:
            k = 0
            for pn, page in enumerate(pdf.pages, 1):
                text = page.extract_text() or ""
                for j, rows in enumerate(page.extract_tables()):
                    t = _to_table(rows, f"Стр. {pn}, таблица {j + 1}", k, header_row=spec.get("header_row"))
                    if t:
                        t.page_text = text
                        t.page = pn
                        tables.append(t)
                        k += 1
    else:
        text = None
        for enc in ("utf-8-sig", "cp1251"):
            try:
                text = data.decode(enc)
                break
            except UnicodeDecodeError:
                continue
        if text is None:
            raise StructureError("Файл не похож на таблицу (не удалось прочитать текст)")
        sample = text[:5000]
        delim = spec.get("delimiter") or max([";", ",", "\t", "|"], key=sample.count)
        rows = list(csv.reader(io.StringIO(text), delimiter=delim))
        t = _to_table(rows, spec.get("title") or "CSV", 0, header_row=spec.get("header_row"))
        if t:
            tables.append(t)
    if not tables:
        raise StructureError("В файле не найдено ни одной таблицы")
    return tables


def pick_file_table(tables, spec):
    """По листу (sheet: имя или номер с 1), странице PDF и тексту над таблицей (after_text) или номеру."""
    spec = spec or {}
    cands = tables
    if spec.get("sheet") not in (None, ""):
        s = spec["sheet"]
        cands = [t for t in tables if str(t.title).lower() == str(s).lower()] or \
                ([tables[int(s) - 1]] if str(s).isdigit() and int(s) <= len(tables) else [])
        if not cands:
            # имя файла в архиве часто с датой выгрузки («…_2026-10-20.csv») — сравниваем без цифр
            def bare(x):
                return re.sub(r"[\d_\-. ]+", "", str(x).lower())
            cands = [t for t in tables if bare(t.title) == bare(s)] or (tables if len(tables) == 1 else [])
    if spec.get("page"):
        cands = [t for t in cands if getattr(t, "page", None) == int(spec["page"])]
    if spec.get("after_text"):
        needle = spec["after_text"].lower()
        hit = [t for t in cands if needle in (getattr(t, "page_text", "") or "").lower()]
        if hit:
            cands = hit
    if spec.get("index") is not None and len(cands) > 1:
        idx = int(spec["index"])
        cands = [t for t in cands if t.index == idx] or cands[idx:idx + 1]
    if not cands:
        raise StructureError(f"В файле не найдена нужная таблица ({spec})")
    return cands[0]


def name_from_url(url) -> str:
    """Название набора по адресу: …/csv/ru/ver-izmenenie-trat.csv → «ver izmenenie trat»;
    …/download/inflyaciya-s-sezonnoj-korrektirovkoj/csv → «inflyaciya s sezonnoj korrektirovkoj»."""
    parts = [p for p in str(url).split("?")[0].split("/") if p]
    skip = {"csv", "zip", "xlsx", "xls", "download", "file", "export", "ru", "en", "data", "api", "v1", "v2"}
    for p in reversed(parts[2:]):
        stem = p.rsplit(".", 1)[0] if "." in p else p
        if stem.lower() not in skip and not stem.isdigit():
            return stem.replace("-", " ").replace("_", " ")
    return ""


def find_link(html_text, base_url, pattern):
    """Ссылка на странице, у которой текст или адрес подходит под шаблон (свежий отчёт)."""
    from urllib.parse import urljoin

    from lxml import html as LH
    doc = LH.fromstring(html_text)
    rx = re.compile(pattern, re.I)
    for a in doc.xpath("//a[@href]"):
        href = a.get("href")
        text = " ".join(a.text_content().split())
        if rx.search(href) or rx.search(text):
            return urljoin(base_url, href)
    raise StructureError(f"На странице не найдена ссылка по шаблону «{pattern}»")


def file_links(doc, base_url):
    from urllib.parse import urljoin
    out = []
    for a in doc.xpath("//a[@href]"):
        href = a.get("href") or ""
        if re.search(r"\.(pdf|xlsx?|csv)(\?|$)", href, re.I):
            out.append((" ".join(a.text_content().split()) or href.rsplit("/", 1)[-1], urljoin(base_url, href)))
    seen, uniq = set(), []
    for t, u in out:
        if u not in seen:
            seen.add(u)
            uniq.append((t, u))
    return uniq
