"""Окно собирается и основные сценарии проходят (нужен дисплей: xvfb-run -a pytest)."""
import json
import os
import time
from pathlib import Path

import pytest

tk = pytest.importorskip("tkinter")
pytestmark = pytest.mark.skipif(not os.environ.get("DISPLAY") and os.name != "nt", reason="нет дисплея")


def pump(app, seconds=3.0, until=None):
    end = time.time() + seconds
    while time.time() < end:
        app.root.update()
        if until and until():
            return True
        time.sleep(0.02)
    return False


@pytest.fixture
def app(workdir, fakes, monkeypatch):
    from hub.gui import app as A
    from hub.core.config import Config
    from hub.core.paths import Paths
    monkeypatch.setattr("hub.core.paths.POINTER", workdir / "pointer.txt")
    Config(Paths(workdir).ensure()).update_settings(first_run_done=True)
    shown = []
    from tkinter import messagebox
    for fn in ("showinfo", "showwarning", "showerror"):
        monkeypatch.setattr(messagebox, fn, lambda *a, _fn=fn, **k: shown.append((_fn, a)))
    for fn in ("askyesno", "askokcancel"):
        monkeypatch.setattr(messagebox, fn, lambda *a, **k: True)
    root = tk.Tk()
    a = A.HubApp(root)
    a.clients = fakes
    a.find.cbr.cbr = fakes["cbr"]
    a.find.fedstat.fs = fakes["fedstat"]
    a.find.fedstat.loaded = True      # каталог ЕМИСС — это 20 МБ с сайта; в тесте не грузим
    a.shown = shown
    pump(a, 0.5)
    yield a
    root.destroy()


def test_all_tabs_build(app):
    for tab in app.nb.tabs():
        app.nb.select(tab)
        pump(app, 0.2)
    assert app.overview.big.cget("text") == "Добавьте первый показатель"


def test_add_flow_model_and_formula(app, fakes):
    from hub.gui.dialog_add import AddDialog
    from hub.gui.tab_jobs import FormulaWindow
    app.cfg.add_model("Сбербанк", str(app.paths.root / "Сбербанк.xlsx"))
    d = AddDialog(app, "cbr", {"category": 18, "indicators": [37], "m1": [2]}, "Ставки по вкладам")
    d.add()
    assert pump(app, 10, until=lambda: "Получено" in d.result.cget("text"))
    assert "1 ряд" in d.result.cget("text") and "Отдать в модель" in d.result.cget("text")
    give = [b for b in d.bar.winfo_children() if "Отдать в модель «Сбербанк»" in str(b.cget("text"))]
    assert give
    give[0].invoke()                                  # окно «Данные для модели» → сохранить выборку целиком
    pump(app, 0.5)
    from hub.gui.dialog_select import SelectionDialog
    sd = next(w for w in app.root.winfo_children() if isinstance(w, SelectionDialog))
    sd.save()
    assert pump(app, 10, until=lambda: app.store.subscribed("Сбербанк") == ["CBR_18_37_m1_2"])
    app.refresh_all()
    assert app.jobs.tree.get_children()
    app.jobs.select(d.job["id"])
    w = FormulaWindow(app, ["CBR_18_37_m1_2"])
    pump(app, 0.2)
    entries = [c for c in w.body.winfo_children() if c.winfo_class() == "TFrame"]
    assert entries
    w.destroy()
    app.models.refresh()
    assert app.models.tree.exists("Сбербанк")
    assert "CBR_18_37_m1_2" in app.models.code_of.values()
    d.destroy()


def test_navigators_with_fake_cbr(app, fakes):
    app.find.cbr.cbr = fakes["cbr"]
    app.nb.select(app.find)              # каталог грузится только на открытой вкладке «Найти данные»
    app.find.nb.select(app.find.cbr)
    assert pump(app, 10, until=lambda: app.find.cbr.roots)
    app.find.cbr.select_indicator(18, 37, "Средневзвешенные процентные ставки по вкладам физических лиц в рублях")
    assert pump(app, 10, until=lambda: str(app.find.cbr.add_btn.cget("state")) == "normal")
    nav = app.find.cbr
    params, name = nav.current_params()
    assert params["category"] == 18 and params["indicators"] == [37] and "y1" not in params
    # годы выбираются списком, как в ЕМИСС: снимаем 2024 → грузим с 2025
    chosen, all_ids, tab, title = nav.lists["y"]
    assert title == "Год" and all_ids == [2024, 2025]
    chosen.discard(2024)
    nav._update_counts()
    assert nav.current_params()[0]["y1"] == 2025
    assert "лет: 1" in nav.sel_var.get()
    assert nav.nb.tab(tab, "text") == "Год (1/2)"
    app.find.q.set("ставки вклад")
    app.find.search()
    kinds = {app.find.results.item(i, "values")[0] for i in app.find.results.get_children()}
    assert "Банк России" in kinds


def test_fedstat_nav_params(app, fakes, tmp_path):
    nav = app.find.fedstat
    nav.fs = fakes["fedstat"]
    nav.cat = {"roots": [], "years": {}}
    nav.run_bg(lambda: nav.fs.indicator("31074"), nav._on_indicator, "…")
    assert pump(app, 20, until=lambda: str(nav.add_btn.cget("state")) == "normal")
    chosen = nav.dim_lists["0"][0]
    chosen.discard("1558900")
    params, name = nav.current_params()
    assert params == {"id": 31074, "select": {"0": ["1558883"]},
                      "select_titles": {"0": ["Российская Федерация"]}}     # название — на случай смены номеров
    assert "Российская Федерация" in name


def test_site_wizard_offline(app, monkeypatch):
    from pathlib import Path
    from hub.sources import _web as W
    web = Path(__file__).resolve().parents[1] / "samples" / "web"
    monkeypatch.setattr(W, "http_get", lambda url, ctx=None, timeout=60: (200, "text/html",
                                                                          (web / "metro_losos.html").read_bytes(), url))
    site = app.find.site
    app.find.nb.select(site)
    site.url.set("https://online.metro-cc.ru/products/losos-murmanskij-6-7kg-inarctica?utm_source=x")
    site.analyze()
    assert pump(app, 15, until=lambda: site.found is not None)
    assert "utm_source" not in site.url.get()
    assert len(site.values.get_children()) == 3
    assert str(site.add_btn.cget("state")) == "normal"
    site.add()
    pump(app, 0.3)
    dialogs = [w for w in app.root.winfo_children() if w.winfo_class() == "Toplevel"]
    assert dialogs and dialogs[-1].name.get().startswith("Лосось")
    dialogs[-1].destroy()


def test_file_tab_preview(app, tmp_path, monkeypatch):
    import datetime as dt
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Нефть"
    ws.append(["Дата", "Брент"])
    ws.append([dt.datetime(2026, 9, 1), 80.0])
    ws.append([dt.datetime(2026, 10, 1), 82.0])
    f = tmp_path / "выгрузка.xlsx"
    wb.save(f)
    from tkinter import filedialog
    monkeypatch.setattr(filedialog, "askopenfilename", lambda **k: str(f))
    tab = app.find.file
    tab.load_inbox()
    pump(app, 0.3)
    assert tab.list.get_children() and tab.mode.get() == "history"
    tab.add()
    pump(app, 0.3)
    assert (app.paths.inbox / "vygruzka" / "выгрузка.xlsx").exists()


def test_settings_examples_and_exports(app, fakes):
    from hub.excel import exports
    app.cfg.add_job("cbr", "Ставки", {"category": 18, "indicators": [37], "m1": [2]})
    r = app.runner()
    r.run(force=True)
    r.close()
    path = exports.export_job(app.cfg, app.store, app.paths, app.cfg.jobs()[0]["id"])
    from openpyxl import load_workbook
    wb = load_workbook(path)
    assert wb.sheetnames == ["Данные", "Паспорт"]
    data = wb["Данные"]
    head = [c.value for c in data[1]]
    assert head == ["Дата", "Период", "Показатель", "Разрез 1", "Периодичность", "Единицы", "Код", "Обновлено",
                    "Значение"]
    assert data["D2"].value.startswith("до 1 года") and data["G2"].value == "CBR_18_37_m1_2"
    assert data["A2"].value >= data["A3"].value          # свежие сверху
    files = exports.update_job_files(app.cfg, app.store, app.paths, app.cfg.jobs())
    assert all(v.endswith(".xlsx") for v in files.values())


def test_selection_dialog_extract(app, fakes):
    """«Данные для модели»: из показателя с 4 рядами отмечаем регион и период — в модель уходит выжимка."""
    from hub.excel import packet
    from hub.gui.dialog_select import SelectionDialog, parse_bound
    assert parse_bound("2022") == "2022-01-01" and parse_bound("2022-02", end=True) == "2022-02-28"
    job = app.cfg.add_job("fedstat", "Зарплата", {"id": 31074, "select": {"0": ["1558883", "1558900"]}})
    r = app.runner()
    r.run(force=True)
    r.close()
    app.cfg.add_model("Сбербанк", str(app.paths.root / "Сбербанк.xlsx"))
    d = SelectionDialog(app, "Сбербанк", job_id=job["id"])
    pump(app, 1)
    region = "Классификатор объектов административно-территориального деления (ОКАТО)"
    assert set(d.fields) == {region, "Периодичность"}
    assert "Подходит рядов: 4" in d.count.cget("text")
    d.fields[region][0].discard("Центральный федеральный округ")
    d.fields["Периодичность"][0].discard("годовая")
    d.d_from.set("2024")
    d.recount()
    assert "Подходит рядов: 1, значений за период: 20" in d.count.cget("text")
    assert "Российская Федерация" in d.name.get() and "с 2024" in d.name.get()
    d.save()
    sel = app.cfg.model("Сбербанк")["selections"][0]
    assert sel["filters"] == {region: ["Российская Федерация"], "Периодичность": ["месячная"]}
    assert sel["from"] == "2024-01-01"
    codes, windows = packet.model_plan(app.cfg, app.store, app.cfg.model("Сбербанк"))
    assert codes == ["FEDSTAT_31074_643_11_M"] and windows[codes[0]] == ("2024-01-01", "")
    app.models.refresh()
    groups = app.models.series.get_children()
    assert len(groups) == 1 and "Выборка" in app.models.series.item(groups[0], "text")


def test_schedule_box_weekly_time_and_release(app):
    from hub.gui.dialog_add import AddDialog
    from hub.gui.schedule_box import KINDS, ReleaseDialog, RuleDialog
    from hub.gui.tab_jobs import EditWindow
    d = AddDialog(app, "cbr", {"category": 18, "indicators": [37], "m1": [2]}, "Ставки")
    box = d.sched_box
    assert box.get() == ("auto", [])
    box.mode.set("rules")
    box._render()                                   # по умолчанию — понедельник
    assert box.rules == [("weekly", (0,), (10, 0))]
    rd = RuleDialog(box, (0, box.rules[0]), box._rule_saved)
    rd.days[0].set(False)
    rd.days[2].set(True)
    rd.time.set("10:30")
    pump(app, 0.1)
    assert "среду в 10:30" in rd.preview.cget("text") and "Ближайшие" in rd.preview.cget("text")
    rd.time.set("25:00")
    pump(app, 0.1)
    assert "Непонятное время" in rd.preview.cget("text")
    rd.time.set("10:30")
    rd.save()
    q = RuleDialog(box, None, box._rule_saved)          # второе правило: 10-й рабочий день после квартала
    q.kind.set(dict(KINDS)["quarterly"])
    q._render()
    q.q_n.set("10")
    q.save()
    assert box.get()[0] == "weekly:Wed@10:30;quarterly:wd10@10:00"
    assert "Ближайшие проверки" in box.info.cget("text")
    rel = ReleaseDialog(box, [], box._set_release)
    rel.text.insert("1.0", "ИПЦ 21.10.2099 19:00\nВВП 18 ноября 2099")
    rel.save()
    assert box.get()[1] == ["2099-10-21T19:00", "2099-11-18"]
    d.add()
    assert pump(app, 10, until=lambda: "Получено" in d.result.cget("text"))
    job = app.cfg.job(d.job["id"])
    assert job["schedule"] == "weekly:Wed@10:30;quarterly:wd10@10:00"
    assert job["release_dates"] == ["2099-10-21T19:00", "2099-11-18"]
    d.destroy()
    w = EditWindow(app, job["id"])
    assert w.sched_box.get() == ("weekly:Wed@10:30;quarterly:wd10@10:00", ["2099-10-21T19:00", "2099-11-18"])
    w.sched_box.lb.selection_set(1)
    w.sched_box.del_rule()
    w.save()
    assert app.cfg.job(job["id"])["schedule"] == "weekly:Wed@10:30"

def test_overview_survives_huge_problem_list(app):
    """Тысячи замечаний (как у большого показателя ЕМИСС) не должны вешать окно."""
    import time as T
    huge = "; ".join(f"Сменились единицы у «ряд {i}»: было «Тонна;^т», стало «Тонна; т»" for i in range(3000))
    for i in range(120):
        j = app.cfg.add_job("cbr", f"Показатель {i}", {"category": 18, "indicators": [37]})
        run = app.store.start_run(j["id"])
        for _ in range(3):
            app.store.new_quarantine(j["id"], run, [huge])
        app.store.set_state(j["id"], status="quarantine", message=huge)
    t = T.time()
    app.overview.refresh()
    pump(app, 0.3)
    assert T.time() - t < 10
    texts = [w.cget("text") for w in app.overview.body.inner.winfo_children() if w.winfo_class() == "TLabel"]
    assert any("ещё 80" in x for x in texts)
    from hub.gui.tab_overview import review_quarantine
    first = app.cfg.jobs()[0]
    review_quarantine(app, app.store.pending_quarantine(first["id"])[-1]["id"], first)
    pump(app, 0.2)


def test_remove_thousands_of_loose_series_fast(app, fakes):
    """«Убрать» у «Отдельных рядов» на тысячи рядов: одна запись файла, ряды выборки не теряются."""
    import time as T
    job = app.cfg.add_job("fedstat", "Зарплата", {"id": 31074, "select": {"0": ["1558883", "1558900"]}})
    r = app.runner()
    r.run(force=True)
    r.close()
    app.cfg.add_model("М", str(app.paths.root / "М.xlsx"))
    codes = [s["code"] for s in app.store.series_of_job(job["id"])]
    app.cfg.save_selection("М", {"job": job["id"], "name": "Выборка", "filters": {}, "from": "", "to": ""})
    models = app.cfg.models()
    models[0]["series"] = codes + [f"OLD_{i}" for i in range(3000)]
    app.cfg.save_models(models)
    app.models.refresh()
    app.models.tree.selection_set("М")
    app.models.show_series()
    loose = next(g for g, k in app.models.groups.items() if k[0] == "codes")
    assert "3004 рядов" in app.models.series.item(loose, "text")
    app.models.series.selection_set(loose)
    t = T.time()
    app.models.remove_series()
    assert T.time() - t < 3
    m = app.cfg.model("М")
    assert m["series"] == [] and m["exclude"] == []
    assert pump(app, 10, until=lambda: not app.tasks)
    assert app.store.subscribed("М") == codes      # ряды выборки на месте



def test_layout_dialog_saves_choice(app):
    from hub.gui.tab_models import LayoutDialog
    app.cfg.add_model("М", str(app.paths.root / "М.xlsx"))
    app.models.refresh()
    d = LayoutDialog(app.models, app.cfg.model("М"))
    assert d.var.get() == "long"
    d.var.set("wide")
    d.apply()
    assert app.cfg.model("М")["layout"] == "wide"
    assert any("Сохранено" in str(a) for _fn, a in app.shown)


def test_dimlist_shift_click_selects_range(app):
    from tkinter import ttk as T
    from hub.gui.dimlist import add_dim_tab
    top = tk.Toplevel(app.root)
    top.geometry("400x400")
    nb = T.Notebook(top)
    nb.pack(fill="both", expand=True)
    entry = add_dim_tab(nb, "Регион", [(str(i), f"Регион {i}") for i in range(10)], lambda: None)
    pump(app, 0.5)
    chosen = entry[0]
    lb = next(w for w in entry[2].winfo_children()[1].winfo_children() if w.winfo_class() == "Listbox")
    chosen.clear()
    lb.selection_clear(0, "end")
    y = lambda i: lb.bbox(i)[1] + 2
    lb.event_generate("<Button-1>", x=5, y=y(2))
    lb.selection_set(2)                       # обычный щелчок отмечает строку
    lb.event_generate("<<ListboxSelect>>")
    lb.event_generate("<Shift-Button-1>", x=5, y=y(6))
    pump(app, 0.1)
    assert chosen == {"2", "3", "4", "5", "6"}


def test_blocked_site_banner_and_site_check(app, monkeypatch):
    import json
    from hub.core import hostblock as HB
    app.cfg.add_job("fedstat", "Зарплата", {"id": 31074, "select": {"0": ["1558883"]}})
    HB.block(app.store, "www.fedstat.ru", "Сайт временно не пускает хаб")
    app.overview.refresh()
    pump(app, 0.2)

    def texts(w):
        out = []
        for c in w.winfo_children():
            if c.winfo_class() in ("TLabel", "TButton"):
                out.append(str(c.cget("text")))
            out += texts(c)
        return out
    all_text = " ".join(texts(app.overview.body.inner))
    assert "Сайт временно не пускает хаб" in all_text and "ЕМИСС (fedstat.ru)" in all_text
    assert "Открыть сайт" in all_text and "Попробовать сейчас" in all_text

    class FakeBrowser:
        def available(self):
            return True

        def login(self, url, banner=None):
            assert url == "https://www.fedstat.ru/"
            return {"cookies": [{"name": "guard", "value": "1", "domain": ".fedstat.ru", "path": "/"}],
                    "user_agent": "Chrome"}
    monkeypatch.setattr(app, "browser", lambda: FakeBrowser())
    ran = []
    monkeypatch.setattr(app, "update_jobs", lambda ids=None, **k: ran.append(ids))
    app.overview.site_check("www.fedstat.ru")
    assert pump(app, 5, until=lambda: ran)
    saved = json.loads((app.paths.cache / "cookies" / "www.fedstat.ru.json").read_text(encoding="utf-8"))
    assert saved["cookies"][0]["name"] == "guard"
    assert HB.until(app.store, "www.fedstat.ru") is None and len(ran[0]) == 1


def test_quarantine_decision_window(app, fakes):
    """Окно отложенных данных: объяснение, кнопки видны, решение в фоне, понятный итог, «не откладывать»."""
    from hub.gui.tab_overview import review_quarantine
    job = app.cfg.add_job("cbr", "Ставки", {"category": 18, "indicators": [37], "m1": [2]},
                          checks={"max_jump_pct": 20})
    r = app.runner()
    r.run(force=True)
    fakes["cbr"].add_month(37, 2, "2025-11-01", 95.0)          # скачок с 19 до 95
    r.run(force=True)
    r.close()
    assert app.store.state(job["id"])["status"] == "quarantine"
    qid = app.store.pending_quarantine(job["id"])[0]["id"]
    review_quarantine(app, qid, app.cfg.job(job["id"]))
    pump(app, 0.3)
    w = next(x for x in app.root.winfo_children() if isinstance(x, tk.Toplevel) and "Отложенные" in x.title())

    def walk(x):
        yield x
        for c in x.winfo_children():
            yield from walk(c)
    texts = " ".join(str(c.cget("text")) for c in walk(w) if c.winfo_class() in ("TLabel", "TButton", "TCheckbutton"))
    assert "Что случилось" in texts and "резко изменилось" in texts and "Новых значений: 1" in texts
    accept = next(c for c in walk(w) if c.winfo_class() == "TButton" and "Принять" in str(c.cget("text")))
    assert accept.winfo_ismapped()                              # кнопка видна
    box = next(c for c in walk(w) if c.winfo_class() == "TCheckbutton")
    box.invoke()                                                # больше не откладывать из-за скачков
    accept.invoke()
    assert pump(app, 10, until=lambda: any("данные приняты" in str(a) for _f, a in app.shown))
    assert app.store.state(job["id"])["status"] == "updated" and not app.store.pending_quarantine(job["id"])
    assert "max_jump_pct" not in (app.cfg.job(job["id"]).get("checks") or {})


def test_change_dataset_in_navigator(app, fakes, monkeypatch):
    """«Изменить набор данных»: навигатор ЕМИСС с отмеченными значениями показателя, «Сохранить изменения»
    меняет параметры того же показателя (id и коды прежние)."""
    job = app.cfg.add_job("fedstat", "Зарплата", {"id": 31074, "select": {"0": ["1558883"]}, "years_from": 2024})
    ran = []
    monkeypatch.setattr(app, "update_jobs", lambda ids=None, **k: ran.append(ids))
    app.jobs.refresh()
    app.jobs.tree.selection_set(job["id"])
    app.jobs.change_data()
    nav = app.find.fedstat
    assert pump(app, 5, until=lambda: "0" in nav.dim_lists and not nav.busy)
    pump(app, 0.3)
    region = nav.dim_lists["0"]
    assert region[0] == {"1558883"}                                  # отмечено как в показателе
    years = nav.dim_lists["3"][0]
    labels = nav.info["filters"]["3"]["values"]
    assert {labels[v].strip()[:4] for v in years} == {y for y in ("2024", "2025") if y in
                                                      {t.strip()[:4] for t in labels.values()}}
    assert "Сохранить" in str(nav.add_btn.cget("text"))
    region[0].add("1558900")                                         # добавили регион
    nav._auto_confirm = False
    nav._add_to_updates()
    new = app.cfg.job(job["id"])
    assert new["params"]["id"] == 31074 and "0" not in new["params"].get("select", {})   # теперь «все регионы»
    assert new["params"].get("years_from") == 2024 and ran == [[job["id"]]]
    assert "В обновления" in str(nav.add_btn.cget("text"))           # режим изменения закончен
    assert len(app.cfg.jobs()) == 1


def test_job_window_schedule_and_history(app, fakes):
    from hub.gui.tab_jobs import JobWindow
    job = app.cfg.add_job("cbr", "Ставки", {"category": 18, "indicators": [37], "m1": [2]})
    r = app.runner()
    r.run(force=True)
    r.close()
    w = JobWindow(app, job["id"])
    pump(app, 0.3)
    assert w.series.get_children() and w.obs.get_children()          # история значений на месте
    w.sched_box.mode.set("rules")
    w.sched_box._render()
    w.save()
    assert app.cfg.job(job["id"])["schedule"] == "weekly:Mon@10:00" and "Сохранено" in w.saved.cget("text")


def test_fedstat_fields_cached_and_failure_saved(app, fakes, monkeypatch):
    """Поля показателя запоминаются на сутки; если не загрузились — начало страницы и ошибка для Claude."""
    nav = app.find.fedstat
    nav.fs = fakes["fedstat"]
    nav.cat = {"roots": [], "years": {}}
    nav._load_indicator("31074")
    assert pump(app, 10, until=lambda: "0" in nav.dim_lists and not nav.busy)
    assert nav._info_cache("31074").exists()
    calls = len(fakes["fedstat"].calls)
    nav._load_indicator("31074")                                   # второй раз — из сохранённого, без сайта
    pump(app, 0.3)
    assert len(fakes["fedstat"].calls) == calls and "0" in nav.dim_lists
    fakes["fedstat"].page = b"<html><body>Showcase without filters</body></html>"
    nav._load_indicator("99999")
    assert pump(app, 10, until=lambda: any("Поля показателя не загрузились" in str(a) for _f, a in app.shown))
    saved = list(app.paths.claude.glob("*emiss_99999*"))
    assert saved and "Showcase without filters" in saved[0].read_text(encoding="utf-8")


def _all_widgets(w):
    out = [w]
    for c in w.winfo_children():
        out += _all_widgets(c)
    return out


def test_overview_news_calendar_filters(app):
    """«Обзор»: новые данные — «по … период, +N записей (дата)», календарь с закрашенным днём и подсказкой,
    кнопки-фильтры источников сохраняются в настройках."""
    import datetime as dt
    from hub.gui.dialog_add import AddDialog
    d = AddDialog(app, "cbr", {"category": 18, "indicators": [37], "m1": [2]}, "Ставки по вкладам")
    d.add()
    assert pump(app, 10, until=lambda: "Получено" in d.result.cget("text"))
    ov = app.overview
    ov.refresh()
    pump(app, 0.3)
    texts = [w.cget("text") for w in _all_widgets(ov.news) if w.winfo_class() in ("TLabel", "Label")]
    line = next(t for t in texts if "запис" in t)
    assert line.startswith("по ") and "+" in line and " в " in line
    today = dt.date.today()
    day = next(w for w in _all_widgets(ov.cal) if w.winfo_class() == "Label" and w.cget("text") == str(today.day))
    assert "Ставки по вкладам" in day.tooltip.text
    # фильтр «Сайты» — показателя ЦБ не видно; настройка сохранилась
    var, btn, _t = ov.groups["sites"]
    btn.invoke()
    pump(app, 0.2)
    assert app.cfg.settings()["overview_sources"] == ["sites"]
    assert "✓ Сайты" == ov.groups["sites"][1].cget("text")
    texts = [w.cget("text") for w in _all_widgets(ov.news) if w.winfo_class() in ("TLabel", "Label")]
    assert not any("запис" in t for t in texts)
    ov.groups["cbr"][1].invoke()
    pump(app, 0.2)
    texts = [w.cget("text") for w in _all_widgets(ov.news) if w.winfo_class() in ("TLabel", "Label")]
    assert any("запис" in t for t in texts)
    ov.refresh()                          # после перерисовки фильтры те же
    assert ov.chosen_groups() == {"sites", "cbr"}
    ov._shift_month(-1)
    ov._shift_month(1)


def test_jobs_sort_by_heading(app):
    from hub.gui import widgets as W
    tree = app.jobs.tree
    for i, (name, upd) in enumerate((("Бета", "05.10.2026 10:00"), ("альфа", "—"), ("Гамма", "21.09.2026 09:00"))):
        tree.insert("", "end", iid=f"x{i}", values=(name, "", "", "", "", "", "", upd))
    assert tree.heading("name", "command")      # заголовок кликабельный
    tree._sort.update(col="name", rev=False)
    W.resort(tree)
    assert [tree.set(i, "name") for i in tree.get_children()][:3] == ["альфа", "Бета", "Гамма"]
    tree._sort.update(col="updated", rev=True)
    W.resort(tree)
    assert [tree.set(i, "updated") for i in tree.get_children()] == ["05.10.2026 10:00", "21.09.2026 09:00", "—"]
    assert W.sort_key_fn(["10 рядов", "9 рядов"])("10 рядов") == 10
    key = W.sort_key_fn(["+5 ч", "−2 дн", "+1 дн", "0"])                 # «Дельта»: часы и дни вместе
    assert sorted(["+5 ч", "−2 дн", "+1 дн", "0"], key=key) == ["−2 дн", "0", "+5 ч", "+1 дн"]
    assert W.sort_key_fn(["2026-01", "2025-Q3"])("2026-01") == "2026-01"


def test_form_dialog_builds_web_form_job(app):
    """Окно «Данные из формы на сайте»: поля периода закрыты (их перебирает хаб), остальные можно поменять;
    «Далее» открывает добавление показателя с параметрами web_form."""
    from hub.gui.dialog_add import AddDialog
    from hub.gui.dialog_form import FormDialog
    from hub.sources import _forms as F
    import customs_site as CS
    from test_forms import REQ
    cand = dict(REQ, text=CS.table({"year": "2025", "startMonth": "1", "endMonth": "12", "pageSize": "10",
                                     "outerValue1": "27111100", "outerValue2": "344"}))
    _c, table = F.best_candidate([cand])
    d = FormDialog(app, "http://stats.customs.gov.cn/indexEn", cand, table)
    pump(app, 0.3)
    assert "year" not in d.vars and "outerValue1" in d.vars
    d.from_year.set("2020")
    d.vars["outerValue2"].set("643")
    pump(app, 0.2)
    assert d.by.get() == "month" and "81 запрос" in d.count.cget("text")     # по месяцу: 2020–сентябрь 2026
    d.by.set("year")
    d._count()
    assert "7 запросов" in d.count.cget("text")
    p = d.params()
    assert p["from_year"] == 2020 and p["values"] == {"outerValue2": "643"} and p["by"] == "year"
    opened = []
    orig = AddDialog.__init__
    AddDialog.__init__ = lambda self, app_, src, params, name, **kw: opened.append((src, params, name))
    try:
        d.next()
    finally:
        AddDialog.__init__ = orig
    assert opened[0][0] == "web_form" and opened[0][2].startswith("Таможня Китая: Liquefied natural gas")
    assert "text" not in opened[0][1]["request"]                    # ответ сайта в настройки не попадает
    # «Изменить набор данных»: тот же показатель, поменяли страну и год
    job = app.cfg.add_job("web_form", "СПГ", opened[0][1])
    e = FormDialog(app, job["params"]["url"], job["params"]["request"], None, job=job)
    pump(app, 0.2)
    assert e.vars["outerValue2"].get() == "643" and e.from_year.get() == "2020"
    e.vars["outerValue2"].set("344")
    e.from_year.set("2022")
    e.save()
    pump(app, 0.3)
    p = app.cfg.job(job["id"])["params"]
    assert p["values"] == {} and p["from_year"] == 2022


def test_comtrade_tab_preview_and_add(app, monkeypatch):
    import json
    from hub.gui.dialog_add import AddDialog
    from hub.sources import _web as Wb, comtrade as CT
    monkeypatch.setattr(CT, "PAUSE_S", 0)

    def fake_get(url, ctx=None, timeout=60):
        from urllib.parse import parse_qs, urlsplit
        pers = parse_qs(urlsplit(url).query)["period"][0].split(",")
        data = [{"refYear": int(p[:4]), "refMonth": int(p[4:]), "netWgt": 1000.0, "primaryValue": 600.0} for p in pers]
        return 200, "application/json", json.dumps({"data": data, "error": ""}).encode(), url
    monkeypatch.setattr(Wb, "http_get", fake_get)
    tab = app.find.comtrade
    tab.preview()
    assert pump(app, 10, until=lambda: "Получено месяцев" in tab.info.cget("text"))
    first = tab.tree.item(tab.tree.get_children()[0])["values"]
    assert str(first[3]) == "600"
    opened = []
    monkeypatch.setattr(AddDialog, "__init__", lambda self, app_, src, params, name, **kw: opened.append((src, params, name)))
    tab.partner.set("123abc")
    assert tab.params() is None                          # непонятная страна — предупреждение
    tab.partner.set("Россия (643)")
    tab.add()
    assert opened == [("comtrade", {"reporter": 156, "partner": 643, "flow": "M", "cmd": "271111", "from_year": 2017},
                       "UN Comtrade: Китай, импорт 271111 — Россия")]


def test_overview_check_row_offers_human_check(app):
    from hub.core import hostblock as HB
    from hub.sources.browser import CHECK_MESSAGE
    HB.block(app.store, "stats.customs.gov.cn", CHECK_MESSAGE)
    app.overview.refresh()
    pump(app, 0.2)
    texts = [w.cget("text") for w in _all_widgets(app.overview.body.inner) if w.winfo_class() in ("TLabel", "TButton")]
    assert "Пройти проверку" in texts and any("ползунок" in t and "Таможня Китая" in t for t in texts)


def test_overview_shows_failed_scheduler_task(app, monkeypatch):
    """Планировщик не может запустить хаб (0x80070002 — файл не найден) — сразу видно на «Обзоре»."""
    from hub.core import winsched
    monkeypatch.setattr(winsched, "status", lambda: (True, "23.10.2026 12:15", 0x80070002))
    app.cfg.add_job("cbr", "Ставки", {"category": 18, "indicators": [37]}, schedule="weekly:Fri@19:05")
    app.sched_fixed = None
    app.overview.refresh()
    pump(app, 0.2)
    texts = [w.cget("text") for w in _all_widgets(app.overview.body.inner) if w.winfo_class() in ("TLabel", "TButton")]
    assert any("не находит файл программы" in t for t in texts) and "Пересоздать задачу" in texts
    assert app.overview.big.cget("text").startswith("Требует внимания")


def test_change_dataset_warns_about_dropped_series_and_models(app, fakes, monkeypatch):
    """Сняли регион в наборе данных: до сохранения — сколько рядов перестанет обновляться и какие модели их
    получают; сохранили — ряды больше не ждём (никаких «исчезли ряды»), по желанию — убраны из моделей."""
    from hub.core import impact
    from hub.gui import nav_edit
    job = app.cfg.add_job("fedstat", "Зарплата", {"id": 31074, "select": {"0": ["1558883", "1558900"]}})
    r = app.runner()
    r.run(force=True)
    r.close()
    series = app.store.series_of_job(job["id"])
    regions = {json.loads(x["dims"])[1][1] for x in series}
    assert len(regions) == 2
    app.cfg.add_model("Модель", str(app.paths.root / "m.xlsx"))
    app.cfg.subscribe("Модель", job_id=job["id"])
    monkeypatch.setattr(app, "update_jobs", lambda ids=None, **k: None)
    app.jobs.refresh()
    app.jobs.tree.selection_set(job["id"])
    app.jobs.change_data()
    nav = app.find.fedstat
    assert pump(app, 5, until=lambda: "0" in nav.dim_lists and not nav.busy)
    pump(app, 0.3)
    nav.dim_lists["0"][0].discard("1558900")                         # сняли второй регион
    seen = {}
    real = nav_edit.confirm_edit

    def spy(parent, job_, dropped, models):
        seen.update(dropped=len(dropped), models=models)
        return True                                                  # «Убрать их из моделей»
    monkeypatch.setattr(nav_edit, "confirm_edit", spy)
    nav._add_to_updates()
    dropped_n = len(series) // 2
    assert seen == {"dropped": dropped_n, "models": [("Модель", dropped_n)]}
    gone = [x["code"] for x in app.store.series_of_job(job["id"]) if x["retired"]]
    assert len(gone) == dropped_n and set(gone) <= set(app.cfg.model("Модель")["exclude"])
    assert app.store.meta_get(f"full:{job['id']}") == "" and app.store.meta_get(f"edited:{job['id']}")
    # следующая (полная) загрузка: снятые ряды не приходят — это не тревога
    r = app.runner()
    out = r.run(job_ids=[job["id"]]).outcomes[0]
    r.close()
    assert out.status != "quarantine", out.problems
    assert not app.store.meta_get(f"edited:{job['id']}")
    # окно подтверждения само по себе собирается и отвечает
    nav._auto_confirm = False
    assert real(nav, job, app.store.series_of_job(job["id"])[:2], [("Модель", 2)]) is False
    assert impact.dropped_series(app.store, app.cfg.job(job["id"]), app.cfg.job(job["id"])["params"], nav.info) == []


def test_app_icon_files_and_window_icon(app):
    from hub.core import setup
    from hub.gui import widgets as W
    assert (W.ASSETS / "hub.png").exists() and W.ICON_ICO.exists() and setup.ICON == W.ICON_ICO
    assert getattr(app.root, "_hub_icon", None) is not None          # значок хаба вместо «пёрышка» Tk
    assert W.ICON_ICO.read_bytes()[:4] == b"\x00\x00\x01\x00"          # настоящий .ico


def test_file_tab_dataset_dialog(app, monkeypatch):
    """«Файл» → длинный CSV: хаб узнаёт набор данных, окно отбора значений, «＋ В обновления» с mode=dataset."""
    from pathlib import Path
    from hub.gui.dialog_add import AddDialog
    from hub.gui.dialog_dataset import DatasetDialog
    from hub.sources import _files as F
    data = (Path(__file__).resolve().parents[1] / "samples" / "web" / "sberindex_dataset.csv").read_bytes()
    url = "https://sberindex.ru/api/dataset/deprecated/csv/ru/ver-izmenenie-trat-po-kategoriyam.csv"
    tab = app.find.file
    tab._show(("url", url), data, F.read_tables(data, "csv"))
    pump(app, 0.3)
    tab.show()
    assert tab.mode.get() == "dataset" and "region (4)" in tab.info.cget("text")
    dialogs = []
    orig_init = DatasetDialog.__init__

    def spy(self, *a, **k):
        orig_init(self, *a, **k)
        dialogs.append(self)
    monkeypatch.setattr(DatasetDialog, "__init__", spy)
    opened = []
    monkeypatch.setattr(AddDialog, "__init__", lambda self, app_, src, params, name, **kw: opened.append((src, params, name)))
    tab.add()
    d = dialogs[0]
    pump(app, 0.4)
    assert "Рядов будет: 16" in d.count.cget("text")
    chosen, _all, _t, _n = d.lists["region"]
    chosen.clear()
    chosen.add("Москва")
    d._count()
    assert "Рядов будет: 4" in d.count.cget("text")
    d.done()
    src, params, name = opened[0]
    assert src == "file_link" and params["mode"] == "dataset" and params["filters"] == {"region": ["Москва"]}
    assert params["url"] == url and name == "Sberindex: ver izmenenie trat po kategoriyam"


def test_fred_link_in_site_tab_opens_fred_tab(app, monkeypatch):
    from hub.gui.dialog_add import AddDialog
    from hub.sources import _web as Wb, fred as FR
    page = (Path(__file__).parent.parent / "samples" / "web")

    def fake_get(url, ctx=None, timeout=60, headers=None):
        if url == FR.csv_url("PFISHUSDM"):
            return 200, "text/csv", (page / "fred_PFISHUSDM.csv").read_bytes(), url
        return 200, "text/html", (page / "fred_PFISHUSDM.html").read_bytes(), url
    monkeypatch.setattr(Wb, "http_get", fake_get)
    monkeypatch.setattr(FR, "_way", {})
    app.find.site.url.set("https://fred.stlouisfed.org/graph/fredgraph.csv?id=PFISHUSDM&cosd=1992-01-01")
    app.find.site.analyze()
    tab = app.find.fred
    assert app.find.nb.select() == str(tab)
    assert pump(app, 10, until=lambda: "1 ряд" in tab.info.cget("text"))
    assert len(tab.tree.get_children()) == 5
    opened = []
    monkeypatch.setattr(AddDialog, "__init__", lambda self, app_, src, params, name, **kw: opened.append((src, params, name)))
    tab.add()
    assert opened[0][0] == "fred" and opened[0][1]["ids"] == ["PFISHUSDM"]
    assert opened[0][2] == "FRED: Global price of Fish meal (PFISHUSDM)"


def test_full_reload_days_in_add_dialog_job_window_and_settings(app, monkeypatch):
    """«Всю историю заново — раз в N дней»: в «＋ В обновления» и ◷ — только у показателей с дозагрузкой."""
    from hub.gui.dialog_add import AddDialog
    from hub.gui.tab_jobs import JobWindow
    d = AddDialog(app, "cbr", {"category": 18, "indicators": [37], "m1": [2]}, "Ставки")
    assert d.reload is None                                     # ЦБ всегда отдаёт всю историю
    d.destroy()
    params = {"url": "https://cbr.ru/currency_base/dynamics/", "request": {"method": "GET", "url": "https://cbr.ru/x"},
              "period": {"kind": "dates", "start": "From", "end": "To", "fmt": "%d.%m.%Y"}, "from_year": 2020}
    d = AddDialog(app, "web_form", params, "Курс евро")
    d.reload.set("abc")
    d.add()
    assert not app.cfg.jobs()                                   # не число — предупреждение, ничего не добавлено
    d.reload.set("0")
    monkeypatch.setattr(app, "bg", lambda *a, **k: None)       # без пробной загрузки (сайта нет)
    d.add()
    job = app.cfg.jobs()[0]
    assert job["full_reload_days"] == 0
    w = JobWindow(app, job["id"])
    pump(app, 0.2)
    assert w.reload.get() == "0"
    w.reload.set("")
    w.save()
    assert app.cfg.job(job["id"])["full_reload_days"] is None and app.cfg.full_reload_days(app.cfg.job(job["id"])) == 30
    tab = app.settings
    tab.refresh()
    assert tab.vars["full_reload_days"].get() == "30"
    tab.vars["full_reload_days"].set("14")
    tab.save()
    assert app.cfg.full_reload_days(app.cfg.job(job["id"])) == 14


def test_form_dialog_vary_values_from_page_list(app):
    """Курсы ЦБ: в окне «Данные страницы» у поля валюты — «☰ Несколько…» со списком валют со страницы; отмеченные
    валюты попадают в params.vary (каждая — свой ряд, за все даты). В «Изменить набор данных» список тот же."""
    import datetime as dt
    import cbr_rates_site as RS
    from hub.gui.dialog_add import AddDialog
    from hub.gui.dialog_form import FormDialog, VaryDialog
    from hub.sources import _forms as F
    page = RS.page(RS.URL, dt.date(2026, 10, 24))
    cand = dict(RS.REQ, resource_type="document", text=page)
    _c, table = F.best_candidate([cand])
    d = FormDialog(app, RS.URL, cand, table, pages=[page])
    pump(app, 0.3)
    field = "UniDbQuery.VAL_NM_RQ"
    assert field in d.choices and field in d.vars and "UniDbQuery.From" not in d.vars
    d.open_vary(field)
    pump(app, 0.3)
    v = [w for w in d.winfo_children() if isinstance(w, VaryDialog)][0]
    assert v.label.get() == "Валюта" and v.values() == [["R01235", "Доллар США"]]   # отмечена валюта из формы
    v.entry[0].add("R01375")
    v.done()
    assert "перебирать 2" in d.vary_hints[field][0].cget("text")
    opened = []
    orig = AddDialog.__init__
    AddDialog.__init__ = lambda self, app_, src, params, name, **kw: opened.append((src, params, name))
    try:
        d.next()
    finally:
        AddDialog.__init__ = orig
    p = opened[0][1]
    assert p["vary"] == [{"field": field, "label": "Валюта", "values": [["R01235", "Доллар США"], ["R01375", "Юань"]]}]
    assert p["vary_choices"][field]["options"][1] == ["R01239", "Евро"] and "Юань" in opened[0][2]
    job = app.cfg.add_job("web_form", "Курсы", p)
    e = FormDialog(app, p["url"], p["request"], None, job=app.cfg.job(job["id"]))
    pump(app, 0.2)
    assert field in e.choices and len(e.vary[field]["values"]) == 2


def test_jobs_list_columns(app, fakes):
    """«Мои показатели»: рядов, период, история, последнее, частота, следующее, модели (и через выборку)."""
    job = app.cfg.add_job("cbr", "Ставки", {"category": 18, "indicators": [37], "m1": [2]},
                          schedule="weekly:Mon,Tue,Wed,Thu,Fri@10:00")
    r = app.runner()
    r.run(force=True)
    r.close()
    app.cfg.add_model("Модель", "x.xlsx")
    app.cfg.subscribe("Модель", job_id=job["id"])
    app.jobs.refresh()
    v = app.jobs.tree.set(job["id"])
    assert v["count"] == "1" and v["period"] == "Месяц" and v["history"] == "2024–2025" and v["last"] == "2025-09"
    assert v["freq"] == "День" and v["next"].isdigit() and v["models"] == "Модель"


def test_future_dates_accepted_once_then_loaded_quietly(app, fakes):
    """Дата из будущего (прогноз): первый раз — отложено; «Принять» ставит «Содержит прогнозные данные», и
    следующие даты в будущем грузятся без вопросов."""
    from hub.gui.tab_overview import review_quarantine
    job = app.cfg.add_job("cbr", "Прогноз", {"category": 18, "indicators": [37], "m1": [2]})
    import datetime as dt
    r = app.runner()
    r.now = dt.datetime(2025, 9, 15, 12, 0)                     # «сегодня» — до последних дат образца
    r.run(force=True)
    fakes["cbr"].add_month(37, 2, "2025-11-01", 19.5)            # октябрь 2025 — пока в будущем
    r.run(force=True)
    assert app.store.state(job["id"])["status"] == "quarantine"
    review_quarantine(app, app.store.pending_quarantine(job["id"])[0]["id"], app.cfg.job(job["id"]))
    pump(app, 0.3)
    w = next(x for x in app.root.winfo_children() if isinstance(x, tk.Toplevel) and "Отложенные" in x.title())

    def walk(x):
        yield x
        for c in x.winfo_children():
            yield from walk(c)
    assert not any(c.winfo_class() == "TCheckbutton" for c in walk(w))      # галочка ставится сама
    next(c for c in walk(w) if c.winfo_class() == "TButton" and "Принять" in str(c.cget("text"))).invoke()
    assert pump(app, 10, until=lambda: any("данные приняты" in str(a) for _f, a in app.shown))
    assert app.cfg.job(job["id"])["checks"].get("future_ok") is True
    fakes["cbr"].add_month(37, 2, "2025-12-01", 19.7)
    rep = r.run(force=True)
    r.close()
    assert rep.outcomes[0].status == "updated"


def test_critical_data_in_model(app, fakes):
    """★ в «Моделях»: о новых данных модели хаб сообщает, только когда обновились критичные ряды."""
    from hub.excel import models as M
    a = app.cfg.add_job("cbr", "Ставки", {"category": 18, "indicators": [37], "m1": [2]})
    b = app.cfg.add_job("cbr", "Масса", {"category": 5, "indicators": [7]})
    app.cfg.add_model("Модель", str(app.paths.root / "нет_такой.xlsx"))
    app.cfg.subscribe("Модель", job_id=a["id"])
    app.cfg.subscribe("Модель", job_id=b["id"])
    r = app.runner()
    r.run(force=True)
    assert dict(M.models_with_new_data(app.cfg, app.store)).get("Модель") == 2
    app.models.refresh()
    tab = app.models
    tab.tree.selection_set("Модель")
    tab.show_series()
    g = next(i for i, k in tab.groups.items() if k == ("job", b["id"]))
    tab.series.selection_set(g)
    tab.toggle_critical()
    assert app.cfg.model("Модель")["critical"] == [f"job:{b['id']}"]
    assert dict(M.models_with_new_data(app.cfg, app.store)).get("Модель") == 1     # только критичный ряд
    assert tab.series.item(next(i for i, k in tab.groups.items() if k == ("job", b["id"])), "text").startswith("★")
    tab.series.selection_set(next(i for i, k in tab.groups.items() if k == ("job", b["id"])))
    tab.toggle_critical()                                       # снять отметку
    assert "critical" not in app.cfg.model("Модель")
    r.close()


def test_columns_dialog_fixes_interpretation(app):
    """«Как понимать таблицу»: ряд можно назвать по-своему; для уже добавленного показателя окно строится по
    последнему ответу сайта, сохранение — «Изменить набор данных» (вся история заново)."""
    import pathlib
    from hub.gui.dialog_add import AddDialog
    from hub.gui.dialog_columns import ColumnsDialog
    from hub.gui.dialog_form import FormDialog
    from hub.sources import _forms as F
    text = (pathlib.Path(__file__).parent.parent / "samples" / "web" / "cbr_avgprocstav.html").read_text("utf-8")
    cand = {"method": "GET", "url": "https://cbr.ru/statistics/avgprocstav/?UniDbQuery.Posted=True",
            "body": "", "content_type": "", "headers": {}, "resource_type": "document", "text": text}
    _c, table = F.best_candidate([cand])
    d = FormDialog(app, "https://cbr.ru/statistics/avgprocstav/", cand, table)
    pump(app, 0.3)
    assert "декадная" in d.summary.cget("text") and "Ставка, %" in d.summary.cget("text")
    d.open_columns()
    pump(app, 0.3)
    c = [w for w in d.winfo_children() if isinstance(w, ColumnsDialog)][0]
    assert "20.10.2026" in str(c.tree.item(c.tree.get_children()[0], "values"))
    c.vars["Ставка, %"][1].set("Макс. ставка по вкладам")
    pump(app, 0.2)
    assert "Макс. ставка по вкладам" in str(c.tree.item(c.tree.get_children()[0], "values"))
    c.done()
    assert "Макс. ставка по вкладам" in d.summary.cget("text")
    opened = []
    orig = AddDialog.__init__
    AddDialog.__init__ = lambda self, app_, src, params, name, **kw: opened.append((src, params, name))
    try:
        d.next()
    finally:
        AddDialog.__init__ = orig
    p = opened[0][1]
    assert p["layout"] == {"names": {"Ставка, %": "Макс. ставка по вкладам"}} and "Макс. ставка" in opened[0][2]
    job = app.cfg.add_job("web_form", "Ставки", p)
    raw = app.paths.raw / job["id"]
    raw.mkdir(parents=True, exist_ok=True)
    (raw / "2026-10-26_120000.html").write_text("<!-- GET x -->\n" + text, encoding="utf-8")
    e = FormDialog(app, p["url"], p["request"], None, job=app.cfg.job(job["id"]))
    pump(app, 0.3)
    assert e.table is not None and "Макс. ставка по вкладам" in e.summary.cget("text")


def test_columns_dialog_for_job_without_saved_answer(app, monkeypatch):
    """У добавленного показателя нет сохранённого ответа (или он обрезан) — «Как понимать таблицу…» всё равно есть:
    хаб спрашивает у сайта один свежий ответ и открывает окно."""
    import pathlib
    from hub.gui.dialog_columns import ColumnsDialog
    from hub.gui.dialog_form import FormDialog
    from hub.sources import web_form as WFm
    text = (pathlib.Path(__file__).parent.parent / "samples" / "web" / "cbr_avgprocstav.html").read_text("utf-8")
    req = {"method": "GET", "url": "https://cbr.ru/statistics/avgprocstav/?UniDbQuery.From=01.12.2024&UniDbQuery.To=02.10.2026",
           "body": "", "content_type": "", "headers": {}, "resource_type": "document"}
    from hub.sources import _forms as F
    params = {"url": "https://cbr.ru/statistics/avgprocstav/", "request": req, "from_year": 2024, "by": "year",
              "period": F.detect_period(F.fields(req)), "periodicity": "auto", "table": {"title": "", "index": 0}}
    job = app.cfg.add_job("web_form", "Ставки", params)
    asked = []
    monkeypatch.setattr(WFm.WebForm, "sample_text", lambda self, p, ctx: asked.append(p) or text)
    e = FormDialog(app, params["url"], req, None, job=app.cfg.job(job["id"]))
    pump(app, 0.3)
    assert e.table is None and "свежий ответ" in e.summary.cget("text")
    e.open_columns()
    assert pump(app, 10, until=lambda: any(isinstance(w, ColumnsDialog) for w in e.winfo_children()))
    assert asked and "декадная" in e.summary.cget("text")


def test_skipped_columns_are_deleted_from_hub(app, monkeypatch):
    """«Не брать» у столбца, который уже загружен: при «Сохранить» хаб предлагает удалить эти ряды совсем —
    вместе со значениями и из моделей, чтобы они нигде не остались."""
    from tkinter import messagebox
    from hub.gui.dialog_form import FormDialog
    from hub.sources import _forms as F
    page = ("<table><tr><th>Дата</th><th>Цена</th><th>Объём</th></tr>" +
            "".join(f"<tr><td>{d:02d}.10.2026</td><td>{100 + d}</td><td>{d * 10}</td></tr>" for d in range(1, 9)) +
            "</table>")
    req = {"method": "GET", "url": "https://example.com/data?From=01.10.2026&To=08.10.2026", "body": "",
           "content_type": "", "headers": {}, "resource_type": "document"}
    params = {"url": "https://example.com/data", "request": req, "from_year": 2026, "by": "year",
              "period": F.detect_period(F.fields(req)), "periodicity": "auto", "table": {"title": "", "index": 0}}
    job = app.cfg.add_job("web_form", "Цены", params)
    r = app.runner()
    r.clients["web_form_http"] = lambda url, ctx=None, timeout=60, headers=None: (200, "text/html", page.encode(), url)
    from hub.sources import web_form as WFm
    monkeypatch.setattr(WFm, "HTTP_PAUSE_S", 0)
    assert r.run(job_ids=[job["id"]]).outcomes[0].status == "updated"
    r.close()
    codes = {s["name"]: s["code"] for s in app.store.series_of_job(job["id"])}
    assert set(codes) == {"Цена", "Объём"}
    app.cfg.add_model("М", "x.xlsx")
    app.cfg.subscribe("М", code=codes["Объём"])
    raw = app.paths.raw / job["id"]
    raw.mkdir(parents=True, exist_ok=True)
    (raw / "1.html").write_text("<!-- GET x -->\n" + page, encoding="utf-8")
    e = FormDialog(app, params["url"], req, None, job=app.cfg.job(job["id"]))
    pump(app, 0.3)
    e._columns_done({"columns": {"Объём": "skip"}}, "auto")
    asked = []
    monkeypatch.setattr(messagebox, "askyesnocancel", lambda *a, **k: asked.append(a) or True)
    monkeypatch.setattr(app, "update_jobs", lambda ids: None)
    e.save()
    assert asked and "Объём" in asked[0][1]
    assert [s["name"] for s in app.store.series_of_job(job["id"])] == ["Цена"]
    assert not app.store.q("SELECT 1 FROM observations WHERE code=?", (codes["Объём"],))
    assert codes["Объём"] not in (app.cfg.model("М").get("series") or [])


def test_sheets_dialog_and_give_to_model(app, monkeypatch):
    """«Файл» → «Листы в модель…»: лист, имя в модели, режим; «＋ В обновления»; «Отдать в модель» подписывает
    модель на листы и предлагает добавить их в книгу."""
    from tkinter import messagebox
    from hub.gui.dialog_add import AddDialog, give_sheets
    from hub.gui.dialog_sheets import SheetsDialog
    data = (Path(__file__).parent.parent / "samples" / "web" / "cbr_bop_est_6.xlsx").read_bytes()
    url = "https://www.cbr.ru/vfs/statistics/credit_statistics/bop/bop_est_6.xlsx"
    d = SheetsDialog(app, {"url": url}, data, "xlsx", "cbr")
    pump(app, 0.3)
    assert d.rows[0][1].get() == "Оценка_ПБ" and "СУММЕСЛИМН(Ячейки_Оценка_ПБ[Значение]" in d.formula.cget("text")
    d.rows[0][1].set("ПБ_оценка")
    d.mode.set("append")
    opened = []
    orig = AddDialog.__init__
    AddDialog.__init__ = lambda self, app_, src, params, name, **kw: opened.append((src, params, name))
    try:
        d.next()
    finally:
        AddDialog.__init__ = orig
    src, params, name = opened[0]
    assert src == "file_sheets" and params["mode"] == "append" and params["url"] == url
    assert params["sheets"] == [{"title": "Оценка ПБ (15.10.2026)", "index": 0, "count": 1, "name": "ПБ_оценка",
                                 "cells": True}]
    job = app.cfg.add_job("file_sheets", name, params)
    app.cfg.add_model("М", str(app.paths.root / "М.xlsx"))
    connected = []
    monkeypatch.setattr(app.models, "_run_connect", lambda n, p: connected.append(n))
    monkeypatch.setattr(messagebox, "askyesno", lambda *a, **k: True)
    # модель заведена после показателя: «Модели» → «Добавить данные» видит и показатель «листы файла»
    from hub.gui.dialog_select import SelectionDialog
    sd = SelectionDialog(app, "М")
    assert job["id"] in sd.job_ids
    sd.jobs_lb.selection_set(sd.job_ids.index(job["id"]))
    sd.pick_job()
    pump(app, 0.3)
    assert "ПБ_оценка (ячейки)" in str([w.cget("text") for w in sd.nb.winfo_children()])
    sd.save()
    assert app.cfg.model("М")["sheets"] == [job["id"]] and connected == ["М"]
    give_sheets(app, "М", job["id"])
    assert app.cfg.model("М")["sheets"] == [job["id"]] and connected == ["М", "М"]
    app.models.refresh()
    app.models.tree.selection_set("М")
    app.models.show_series()
    assert any(k == ("sheets", job["id"]) for k in app.models.groups.values())


def test_path_dialog_check_and_file_tab(app, monkeypatch):
    """«Путь к файлу»: шаги из записанных щелчков → «Проверить путь» (ход по шагам) → файл во вкладке «Файл» →
    «▤ Листы в модель…» хранит путь, а не адрес файла; ✎ у показателя открывает те же шаги."""
    import minfin_site as MS
    from hub.gui.dialog_add import AddDialog
    from hub.gui.dialog_path import PathDialog
    from hub.sources import _path as PT
    from hub.sources import _web as WEB
    base = "https://minfin.gov.ru"
    site = MS.Site()

    def fake_get(url, ctx=None, timeout=60, headers=None):
        path = url.replace(base, "")
        if path.startswith("/common/upload/library/"):
            return 200, "application/vnd.ms-excel", MS.xlsx(path.rsplit("/", 1)[-1]), url
        if path.startswith("/ru/statistics/fedbud"):
            return 200, "text/html; charset=utf-8", MS.list_page(site.releases).encode(), url
        key = path.split("id_4=")[-1]
        return 200, "text/html; charset=utf-8", MS.doc_page(key[0], key[1:]).encode(), url
    monkeypatch.setattr(WEB, "http_get", fake_get)
    monkeypatch.setattr(app, "bg", lambda work, done=None, msg="", error=None: done(work()) if done else work())
    html = MS.list_page(site.releases)
    icon = [x for x in PT.links(html, base + "/ru/statistics/fedbud/") if x["href"].endswith("m2026-10-09.xlsx")][0]
    clicks = [{**icon, "html": html, "page": base + "/ru/statistics/fedbud/"}]
    tab = app.find.file
    d = PathDialog(app, base + "/ru/statistics/fedbud/", PT.guess_steps(clicks), clicks, on_done=tab._path_loaded)
    pump(app, 0.2)
    assert d.rows[0][1]["near"].get() == "краткая ежемесячная информация исполнении"
    d.check()
    out = d.out.get("1.0", "end")
    assert "подходящих ссылок: 1" in out and "m2026-10-09.xlsx" in out
    d.rows[0][1]["near"].set("нет таких слов")
    d.check()
    assert "Не получилось" in d.out.get("1.0", "end")
    d.rows[0][1]["near"].set("ежемесячная")
    d.finish()
    pump(app, 0.2)
    assert tab.source[0] == "path" and tab.file_url.endswith("m2026-10-09.xlsx") and tab.tables
    opened = []
    orig = AddDialog.__init__
    AddDialog.__init__ = lambda self, app_, src, params, name, **kw: opened.append((src, params, name))
    try:
        tab.to_sheets()
        sd = [w for w in app.root.winfo_children() if w.__class__.__name__ == "SheetsDialog"][-1]
        sd.next()
        tab.list.selection_set("0")
        tab.add()
    finally:
        AddDialog.__init__ = orig
    (src1, p1, _n1), (src2, p2, name2) = opened
    assert src1 == "file_sheets" and p1["path"]["steps"][0]["near"] == "ежемесячная" and "url" not in p1
    assert src2 == "file_link" and p2["path"]["start"].endswith("/fedbud/") and "url" not in p2
    assert name2.startswith("Minfin")
    job = app.cfg.add_job("file_link", name2, p2)
    app.jobs.refresh()
    monkeypatch.setattr(app.jobs, "selected", lambda: job["id"])
    app.jobs.change_data()
    e = [w for w in app.root.winfo_children() if isinstance(w, PathDialog)][-1]
    e.rows[0][1]["words"].set("документ")
    e.rows[0][1]["href"].set("")
    e._add_row({"words": "скачать"}, None)
    from tkinter import messagebox
    monkeypatch.setattr(messagebox, "askyesno", lambda *a, **k: False)
    e.save()
    steps = app.cfg.job(job["id"])["params"]["path"]["steps"]
    assert [s["words"] for s in steps] == ["документ", "скачать"]
    assert PT.follow(app.cfg.job(job["id"])["params"]["path"]).endswith("m2026-10-09.xlsx")
