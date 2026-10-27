"""«Листы файла в модель»: лист Excel как есть на постоянное имя листа, режимы «заменять» и «копить обновления»,
таблица ячеек для формул (путь строки по отступам, заголовок столбца по этажам шапки). Образец — файл ЦБ
«Оценка ключевых агрегатов платежного баланса» (bop_est_6.xlsx)."""
import csv
import datetime as dt
import fnmatch
import io
from pathlib import Path

import pytest

from hub.core import sheets as SH
from hub.excel import packet

SAMPLE = Path(__file__).parent.parent / "samples" / "web" / "cbr_bop_est_6.xlsx"
URL = "https://www.cbr.ru/vfs/statistics/credit_statistics/bop/bop_est_6.xlsx"


def next_release() -> bytes:
    """Следующая версия файла: лист с новой датой в имени, шапка сдвинулась (январь–сентябрь, добавлен сентябрь
    вместо июля), значения августа уточнены."""
    from openpyxl import load_workbook
    wb = load_workbook(SAMPLE)
    ws = wb.worksheets[0]
    ws.title = "Оценка ПБ (12.11.2026)"
    ws["B4"] = "январь - сентябрь\n2026 г.\n(оценка)"
    ws["E4"] = "август\n2026 г.\n(оценка)"
    ws["F4"] = "сентябрь\n2026 г.\n(оценка)"
    ws["E7"] = 11.9                    # товары, август — уточнение (было 11,7 в столбце F)
    ws["F7"] = 12.5                    # товары, сентябрь
    ws["A30"] = "Дата последнего обновления: 12 ноября 2026 года."
    ws["A28"] = None
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_cells_have_paths_and_headers():
    sh = SH.read_sheets(SAMPLE.read_bytes(), "xlsx")[0]
    assert sh.title == "Оценка ПБ (15.10.2026)" and SH.file_date(sh) == "2026-10-15"
    assert SH.default_name(sh.title) == "Оценка_ПБ"
    cs = SH.cells(sh)
    goods = [c for c in cs if c["Путь"] == "Счет текущих операций / Товары / Экспорт"]
    services = [c for c in cs if c["Путь"] == "Счет текущих операций / Услуги / Экспорт"]
    assert goods and services and goods[0]["Значение"] == 312.4 and services[0]["Значение"] == 31.1
    heads = {c["Заголовок"] for c in cs}
    assert "август 2026 г. (оценка)" in heads and "Справочно: — январь - август 2025 г." in heads
    assert "(млрд долларов США)" in cs[0]["Таблица"]
    assert SH.valid_name("ПБ_оценка") is None and SH.valid_name("Хаб_X") and SH.valid_name("a" * 30)


def sumifs_latest(rows, path, head):
    """Как формула СУММЕСЛИМН(…; Обновлено; МАКСЕСЛИ(…)) — последнее известное значение (с подстановкой «*»)."""
    hit = [r for r in rows if fnmatch.fnmatchcase(r["Путь"], path) and fnmatch.fnmatchcase(r["Заголовок"], head)]
    last = max(r["Обновлено"] for r in hit)
    return sum(float(r["Значение"]) for r in hit if r["Обновлено"] == last)


@pytest.fixture
def site(monkeypatch):
    from hub.sources import _web as W
    state = {"data": SAMPLE.read_bytes(), "calls": 0}

    def fake_get(url, ctx=None, timeout=60, headers=None):
        state["calls"] += 1
        return 200, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", state["data"], url
    monkeypatch.setattr(W, "http_get", fake_get)
    return state


def test_append_mode_keeps_versions_and_formula_finds_latest(runner, site):
    params = {"url": URL, "mode": "append",
              "sheets": [{"title": "Оценка ПБ (15.10.2026)", "index": 0, "count": 1, "name": "ПБ_оценка", "cells": True}]}
    job = runner.cfg.add_job("file_sheets", "ЦБ: оценка платёжного баланса", params)
    runner.cfg.add_model("М", "x.xlsx")
    runner.cfg.subscribe_sheets("М", job["id"])
    runner.now = dt.datetime(2026, 10, 16, 10, 0)
    o = runner.run(job_ids=[job["id"]]).outcomes[0]
    assert o.status == "updated" and o.changed and "ПБ_оценка" in o.message
    assert runner.store.state(job["id"])["last_label"] == "файл от 15.10.2026"
    o = runner.run(job_ids=[job["id"]]).outcomes[0]                  # тот же файл — ничего не дописано
    assert o.status == "unchanged"
    site["data"] = next_release()
    o = runner.run(job_ids=[job["id"]]).outcomes[0]
    assert o.status == "updated" and "файл от 12.11.2026" in o.message
    pp = packet.packet_path(runner.paths, "М")
    grid = list(csv.reader(io.StringIO(packet.sheet_model_path(pp, "ПБ_оценка").read_text("utf-8-sig")),
                           delimiter=";"))
    assert grid[0][:4] == ["Обновлено", "Файл от", "Строка", "A"]
    assert len(grid) == 1 + 28 + 30 and grid[1][1] == "2026-10-15" and grid[-1][1] == "2026-11-12"
    cells = list(csv.DictReader(io.StringIO(packet.sheet_model_path(pp, "ПБ_оценка", True).read_text("utf-8-sig")),
                                delimiter=";"))
    # август: был в столбце F (11,7), в новой версии — в E (11,9, уточнён); формула берёт последнее
    assert sumifs_latest(cells, "*Товары", "август 2026*") == 11.9
    assert sumifs_latest(cells, "*Товары", "июль 2026*") == 13.3          # июля в новой версии нет — прежнее
    assert sumifs_latest(cells, "*Товары", "сентябрь 2026*") == 12.5
    upd = packet.updates_path(pp).read_text("utf-8-sig")
    assert upd.count("добавлено обновление листа") == 2 and "лист «ПБ_оценка»" in upd
    from hub.excel.models import models_with_new_data
    assert dict(models_with_new_data(runner.cfg, runner.store)).get("М") == 1


def test_replace_mode_and_placeholder_before_first_load(runner, site):
    params = {"url": URL, "mode": "replace",
              "sheets": [{"title": "Оценка ПБ (15.10.2026)", "index": 0, "count": 1, "name": "ПБ", "cells": False}]}
    job = runner.cfg.add_job("file_sheets", "ЦБ: ПБ", params)
    runner.cfg.add_model("М", "x.xlsx")
    runner.cfg.subscribe_sheets("М", job["id"])
    pp = packet.packet_path(runner.paths, "М")
    packet.rebuild(runner.cfg, runner.store, runner.paths)
    assert packet.sheet_model_path(pp, "ПБ").read_text("utf-8-sig").strip() == "Строка;A"   # ещё не загружен
    assert not packet.sheet_model_path(pp, "ПБ", True).exists()
    runner.run(job_ids=[job["id"]])
    site["data"] = next_release()
    runner.run(job_ids=[job["id"]])
    grid = list(csv.reader(io.StringIO(packet.sheet_model_path(pp, "ПБ").read_text("utf-8-sig")), delimiter=";"))
    assert grid[0][:3] == ["Строка", "A", "B"] and len(grid) == 1 + 30
    assert grid[4][2].startswith("январь - сентябрь")                    # всегда последняя версия
    assert grid[7][5] == "11.9"                                          # E7 — число с точкой
