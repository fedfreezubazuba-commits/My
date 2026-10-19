"""Веб-источники на сохранённых страницах samples/web (Metro, Металл Торг, Fearnleys)."""
import datetime as dt
import io
import os
from pathlib import Path

import pytest

from hub.sources import Context, _analyze as A, _web as W, get
from hub.sources import _files as F

WEB = Path(__file__).resolve().parents[1] / "samples" / "web"
METRO_URL = "https://online.metro-cc.ru/products/losos-murmanskij-6-7kg-inarctica?utm_source=yandex&yclid=123&x=1"
NOW = dt.datetime(2026, 10, 19, 12, 0)


def page(name):
    data = (WEB / name).read_bytes()
    return W.decode(data)


@pytest.fixture
def offline(monkeypatch):
    """Подменяем загрузку: вместо сайта — сохранённая страница."""
    pages = {}

    def fake_get(url, ctx=None, timeout=60):
        for key, name in pages.items():
            if key in url:
                data = (WEB / name).read_bytes()
                return 200, "text/html", data, url
        raise AssertionError(url)
    monkeypatch.setattr(W, "http_get", fake_get)
    return pages


def test_clean_url():
    assert W.clean_url(METRO_URL) == "https://online.metro-cc.ru/products/losos-murmanskij-6-7kg-inarctica?x=1"
    assert W.clean_url("metaltorg.ru/#top") == "https://metaltorg.ru/"
    assert W.site_name("https://online.metro-cc.ru/x") == "metro"


def test_metro_wizard_and_fetch(offline):
    offline["metro-cc.ru"] = "metro_losos.html"
    found = A.analyze(page("metro_losos.html"), METRO_URL)
    vals = {v["field"]["key"]: v["value"] for v in found.values}
    assert vals == {"price_piece": 12863.5, "price_kg": 1979.0, "available": 0.0}
    params, name = A.value_params(found)
    assert params["url"].endswith("inarctica?x=1") and "utm_source" not in params["url"]
    assert params["code_base"] == "METRO_losos_murmanskij"
    assert name.startswith("Лосось Мурманск")
    res = get("web_value").fetch(params, Context(now=NOW))
    codes = {s.code: o.value for s, o in zip(res.series, res.obs)}
    assert codes == {"METRO_losos_murmanskij_rub_sht": 12863.5, "METRO_losos_murmanskij_rub_kg": 1979.0,
                     "METRO_losos_murmanskij_nalichie": 0.0}
    assert all(o.period_end == NOW.date() for o in res.obs)
    assert not [w for w in res.warnings if "Расхождение" in w]     # запасные стратегии согласны


def test_web_value_fallback_and_discrepancy(offline):
    offline["metro-cc.ru"] = "metro_losos.html"
    params = {"url": METRO_URL, "fields": [
        {"key": "price_piece", "name": "шт", "strategies": [{"type": "css", "selector": "#нет-такого"},
                                                           {"type": "jsonld", "path": "offers.price"}]},
        {"key": "price_kg", "name": "кг", "strategies": [{"type": "profile", "name": "metro", "field": "price_kg"},
                                                         {"type": "jsonld", "path": "offers.price"}]}]}
    res = get("web_value").fetch(params, Context(now=NOW))
    assert [o.value for o in res.obs] == [12863.5, 1979.0]
    assert any("Расхождение" in w for w in res.warnings)


def test_web_value_structure_error(offline):
    from hub.core.errors import StructureError
    offline["metro-cc.ru"] = "metaltorg_main.html"         # «сайт поменял вёрстку»
    params = {"url": METRO_URL, "fields": [{"key": "price", "name": "x",
                                            "strategies": [{"type": "jsonld", "path": "offers.price"}]}]}
    with pytest.raises(StructureError):
        get("web_value").fetch(params, Context(now=NOW))


def test_metaltorg_snapshot_tables(offline):
    offline["metaltorg.ru"] = "metaltorg_main.html"
    found = A.analyze(page("metaltorg_main.html"), "https://www.metaltorg.ru/")
    titles = [t.title for t in found.tables]
    nonferrous = next(t for t in found.tables if t.title.startswith("Цветные металлы"))
    params, name = A.table_params(found, nonferrous)
    assert name.startswith("Metaltorg: Цветные металлы")
    res = get("web_table").fetch(params, Context(now=NOW))
    by = {s.name.split(" | ")[-1]: (s, o) for s, o in zip(res.series, res.obs)}
    assert set(by) == {"Al", "Cu", "Pb", "Ni", "Sn", "Zn"}
    s, o = by["Cu"]
    assert o.value == 14436 and o.period_end == dt.date(2026, 10, 19) and s.unit == "US$/т"
    assert s.code.startswith("METALTORG_") and len(s.code) <= 40
    steel = next(t for t in found.tables if "сталь" in t.title)
    res = get("web_table").fetch(A.table_params(found, steel)[0], Context(now=NOW))
    assert {o.value for o in res.obs} == {485.7, 460.2}
    assert any("Драгоценные" in t for t in titles)


def test_fearnleys_cards_weekly(offline):
    offline["fearnpulse.com"] = "fearnleys_weekly.html"
    url = "https://fearnpulse.com/fearnleys-weekly-report"
    found = A.analyze(page("fearnleys_weekly.html"), url)
    assert found.page_date == dt.date(2026, 10, 14) and found.page_is_week
    dirty = next(t for t in found.tables if t.kind == "cards" and t.title.startswith("Dirty"))
    params, _ = A.table_params(found, dirty)
    res = get("web_table").fetch(params, Context(now=NOW))
    vals = {s.name.split(" | ")[-1]: o for s, o in zip(res.series, res.obs)}
    assert vals["MEG/WEST 280'"].value == 725
    assert vals["Caribs/USG 70'"].value == 915                 # изменение -135 не путается со ставкой
    assert vals["MEG/WEST 280'"].period_end == dt.date(2026, 10, 14)
    assert vals["MEG/WEST 280'"].period_label == "2026-W42"
    cape = next(t for t in found.tables if t.title.startswith("Capesize"))
    res = get("web_table").fetch(A.table_params(found, cape)[0], Context(now=NOW))
    assert {o.value for o in res.obs} == {71778, 12.05, 28841}
    sale = next(t for t in found.tables if t.kind == "html" and t.header and t.header[0] == "Dry")
    res = get("web_table").fetch(A.table_params(found, sale)[0], Context(now=NOW))
    names = {s.name: o.value for s, o in zip(res.series, res.obs)}
    assert names["Capesize | 5 yr old"] == 75.0 and names["Capesize | 10 yr old"] == 60.0


def test_table_found_by_title_after_reorder(offline):
    offline["metaltorg.ru"] = "metaltorg_main.html"
    params = {"url": "https://www.metaltorg.ru/", "table": {"kind": "html", "title": "Драгоценные металлы", "index": 0}}
    res = get("web_table").fetch(params, Context(now=NOW))
    assert {s.name.split(" | ")[-1] for s in res.series} == {"Золото", "Серебро", "Платина", "Палладий"}


def test_block_detection():
    from hub.core import errors as E
    assert W.detect_block(403, "") == E.BOT
    assert W.detect_block(200, "<html>Подтвердите, что вы не робот. SmartCaptcha</html>") == E.BOT
    assert W.detect_block(429, "slow down") == E.RATE_LIMIT
    assert W.detect_block(200, "<html>ok</html>") is None


def test_js_page_without_browser_explains(monkeypatch):
    """Страница строится скриптом, встроенного браузера нет — понятное объяснение."""
    from hub.core.errors import StructureError
    monkeypatch.setattr(W, "http_get", lambda url, ctx=None, timeout=60: (200, "text/html", b"<html><body>"
                        b"<div id=__next></div></body></html>", url))
    params = {"url": "https://fearnpulse.com/fearnleys-weekly-report", "table": {"kind": "cards", "title": "Dirty"}}
    with pytest.raises(StructureError) as e:
        get("web_table").fetch(params, Context(now=NOW))
    assert "встроенный браузер" in str(e.value)


def _xlsx_bytes():
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Цены"
    ws.append(["Отчёт о ценах"])
    ws.append(["Дата", "Брент", "Юралс"])
    for i, (b, u) in enumerate([(80.1, 60.2), (81.5, 61.0), (79.9, 59.4)]):
        ws.append([dt.datetime(2026, 7 + i, 1), b, u])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_file_link_xlsx_history(monkeypatch):
    data = _xlsx_bytes()
    monkeypatch.setattr(W, "http_get", lambda url, ctx=None, timeout=60: (200, "application/octet-stream", data, url))
    params = {"url": "https://example.ru/prices.xlsx?utm_medium=x", "table": {"sheet": "Цены"}, "periodicity": "month"}
    res = get("file_link").fetch(params, Context(now=NOW))
    names = {s.name for s in res.series}
    assert names == {"Цены | Брент", "Цены | Юралс"}
    brent = sorted((o for o in res.obs if o.key == "col|Брент"), key=lambda o: o.period_end)
    assert [(o.period_end, o.value) for o in brent][0] == (dt.date(2026, 7, 31), 80.1)
    assert brent[0].period_label == "2026-07"


def test_file_link_csv_and_link_pattern(monkeypatch):
    csv_bytes = "Дата;Ставка\n2026-10-01;15,5\n2026-10-08;15,25\n".encode("cp1251")
    html = b'<html><body><a href="/files/old.csv">old</a><a href="/files/weekly_2026_42.csv">Weekly report</a></body></html>'

    def fake(url, ctx=None, timeout=60):
        return (200, "text/html", html, url) if url.endswith("/reports") else (200, "text/csv", csv_bytes, url)
    monkeypatch.setattr(W, "http_get", fake)
    params = {"page_url": "https://example.ru/reports", "link_pattern": r"weekly.*\.csv", "periodicity": "day"}
    res = get("file_link").fetch(params, Context(now=NOW))
    assert res.url == "https://example.ru/files/weekly_2026_42.csv"
    assert sorted(o.value for o in res.obs) == [15.25, 15.5]


def test_pdf_tables():
    pytest.importorskip("pdfplumber")
    pdf = _make_pdf()
    if pdf is None:
        pytest.skip("нет генератора PDF")
    tables = F.read_tables(pdf, "pdf")
    assert tables and any("Capesize" in " ".join(r) for t in tables for r in t.rows)


def _make_pdf():
    """Маленький PDF с таблицей (рисуем линии, чтобы pdfplumber её увидел)."""
    try:
        from reportlab.lib.pagesizes import A4
        from reportlab.platypus import SimpleDocTemplate, Table, TableStyle
    except ImportError:
        return None
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4)
    t = Table([["Vessel", "Rate"], ["Capesize", "71778"], ["Panamax", "21845"]])
    t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.5, (0, 0, 0))]))
    doc.build([t])
    return buf.getvalue()


def test_folder_inbox(workdir, monkeypatch):
    from hub.core.paths import Paths
    from hub.runner import Runner
    paths = Paths(workdir).ensure()
    r = Runner(paths)
    try:
        job = r.cfg.add_job("folder_inbox", "Терминал: нефть", {"folder": "нефть", "table": {"sheet": "Цены"},
                                                               "periodicity": "month"})
        rep = r.run(job_ids=[job["id"]])
        assert rep.outcomes[0].status in ("waiting", "unchanged", "delayed") or rep.outcomes[0].category == "no_new_data"
        inbox = paths.inbox / "нефть"
        (inbox / "выгрузка.xlsx").write_bytes(_xlsx_bytes())
        rep = r.run(job_ids=[job["id"]])
        o = rep.outcomes[0]
        assert o.status == "updated" and len(o.new_periods) == 6
        assert not (inbox / "выгрузка.xlsx").exists()
        assert any(p.name.endswith("выгрузка.xlsx") for p in (inbox / "обработано").iterdir())
        rep = r.run(job_ids=[job["id"]])
        assert not rep.outcomes[0].changed
    finally:
        r.close()


@pytest.mark.skipif(not os.environ.get("HUB_CHROMIUM"), reason="нужен Chromium (HUB_CHROMIUM)")
def test_browser_picker_on_sample(tmp_path):
    """«Выбрать на странице»: щелчок по цене в окне браузера → селектор, по которому хаб найдёт число."""
    pytest.importorskip("playwright")
    from hub.core.paths import Paths
    from hub.sources.browser import Browser
    br = Browser(Paths(tmp_path))
    html = tmp_path / "p.html"
    html.write_text("<html><body><h1>Товар</h1><div class='card'><span class='lbl'>Цена за кг</span>"
                    "<span class='price'>1 979,50 ₽</span></div><table><tr><td>Al</td><td>3001</td></tr>"
                    "<tr><td>Cu</td><td>14436</td></tr></table></body></html>", "utf-8")
    pick = br.pick(html.as_uri(), headless=True, timeout_s=20, on_page=lambda p: p.click("span.price"))
    assert pick and pick["kind"] == "value" and "1 979,50" in pick["text"]
    params, label = A.picked_value_params("https://shop.example/x", pick)
    assert label == "Цена за кг"
    from hub.sources.web_value import extract_fields
    doc = W.parse_html(pick["page_html"])
    got = extract_fields(params["fields"], doc, pick["page_html"], "x")
    assert got["value"][0] == 1979.5
    pick = br.pick(html.as_uri(), headless=True, timeout_s=20,
                   on_page=lambda p: p.click("td >> nth=1", modifiers=["Shift"]))
    assert pick["kind"] == "table"
    t = A.picked_table("x", pick)
    assert t is not None and t.rows[1] == ["Cu", "14436"]
    html_text = br.get_html(html.as_uri(), wait_ms=100)[0]
    assert "14436" in html_text


def test_example_jobs_work_on_samples(offline, workdir):
    """Показатели-образцы (Metro, Металл Торг, Fearnleys) — через весь конвейер до посылки."""
    from hub.core.examples import add_examples
    from hub.core.paths import Paths
    from hub.excel import packet
    from hub.runner import Runner
    offline.update({"metro-cc.ru": "metro_losos.html", "metaltorg.ru": "metaltorg_main.html",
                    "fearnpulse.com": "fearnleys_weekly.html"})
    r = Runner(Paths(workdir), now=NOW)
    try:
        ids = add_examples(r.cfg)
        assert len(ids) == 5 and add_examples(r.cfg) == []
        r.cfg.add_model("Металлурги", "x.xlsx")
        for i in ids:
            r.cfg.subscribe("Металлурги", job_id=i)
        rep = r.run(force=True)
        bad = [(o.job["name"], o.message) for o in rep.outcomes if o.status != "updated"]
        assert not bad, bad
        codes = r.store.subscribed("Металлурги")
        assert "METRO_losos_murmanskij_rub_kg" in codes
        assert all(len(c) <= 40 and c.replace("_", "").isalnum() for c in codes)
        rows = packet.packet_path(r.paths, "Металлурги").read_text("utf-8-sig").splitlines()
        assert any(";METRO_losos_murmanskij_rub_kg;" in x and x.endswith(";1979") for x in rows)
        assert any(x.startswith("2026-10-14;2026-W42;FEARNLEYS_") for x in rows)
        # повторный прогон в тот же день ничего не меняет
        rep = r.run(force=True)
        assert not any(o.changed for o in rep.outcomes)
    finally:
        r.close()


def test_snapshot_history_builds_without_false_alarms(offline, workdir):
    """Две недели ежедневных снимков: история копится, карантина нет, пересмотров нет."""
    from hub.core.paths import Paths
    from hub.runner import Runner
    offline["metaltorg.ru"] = "metaltorg_main.html"
    paths = Paths(workdir)
    job = None
    for day in range(14):
        r = Runner(paths, now=NOW + dt.timedelta(days=day))
        try:
            if job is None:
                job = r.cfg.add_job("web_table", "Металлы", {"url": "https://www.metaltorg.ru/", "date": "today",
                                                             "table": {"kind": "html", "title": "Цветные металлы"}},
                                    checks={"max_jump_pct": 30})
            rep = r.run(job_ids=[job["id"]])
            assert rep.outcomes[0].status == "updated", rep.outcomes[0].message
            assert not rep.outcomes[0].revisions
        finally:
            r.close()
    r = Runner(paths)
    try:
        code = r.store.series_of_job(job["id"])[0]["code"]
        assert len(r.store.current_rows(code)) == 14
    finally:
        r.close()


def test_scenario_layout_changed_red_status_and_claude(offline, workdir):
    """Сценарий 4 ТЗ: сайт поменял вёрстку → красный статус на «Обзоре» → «Отправить Claude» → текст."""
    from hub.core import claude_pack, overview
    from hub.core.paths import Paths
    from hub.runner import Runner
    from hub.sources import _analyze
    offline["metro-cc.ru"] = "metro_losos.html"
    r = Runner(Paths(workdir), now=NOW)
    try:
        params, name = _analyze.value_params(_analyze.analyze(page("metro_losos.html"), METRO_URL))
        job = r.cfg.add_job("web_value", name, params)
        assert r.run(job_ids=[job["id"]]).outcomes[0].status == "updated"
        offline["metro-cc.ru"] = "metaltorg_main.html"           # «новая вёрстка»
        out = r.run(job_ids=[job["id"]]).outcomes[0]
        assert out.status == "error" and out.category == "structure"
        probs = overview.problems(r.cfg, r.store)
        assert probs and probs[0].action == "claude" and probs[0].action_label == "Отправить Claude"
        text = claude_pack.job_package(r.cfg, r.store, job["id"])
        assert "web_value" in text and "вёрстка" in text and "METRO_losos_murmanskij_rub_kg" in text
        assert "Начало последнего сырого ответа" in text
    finally:
        r.close()


@pytest.mark.skipif(not os.environ.get("HUB_CHROMIUM"), reason="нужен Chromium (HUB_CHROMIUM)")
def test_pick_block_on_fearnleys_becomes_table(tmp_path):
    """Щелчок по блоку «плиток» Fearnleys (как сделал пользователь) — это выбор таблицы, а не одного числа."""
    pytest.importorskip("playwright")
    from hub.core.paths import Paths
    from hub.sources.browser import Browser
    br = Browser(Paths(tmp_path))
    url = (WEB / "fearnleys_weekly.html").as_uri()

    def click_grid(p):
        p.locator("text=MEG/WEST").first.evaluate(
            "el => el.closest('.grid').dispatchEvent(new MouseEvent('click', {bubbles: true}))")
    pick = br.pick(url, headless=True, timeout_s=30, on_page=click_grid)
    assert pick and pick["kind"] == "table"
    assert "click" not in (pick["label"] or "").lower()
    t = A.picked_table("https://fearnpulse.com/fearnleys-weekly-report", pick)
    assert t is not None and t.title.startswith("Dirty")


def test_picked_table_by_text_when_html_missing():
    html = page("fearnleys_weekly.html")
    labels = ["MEG/WEST", "MEG/Japan", "MEG/Singapore", "WAF/FEAST", "WAF/USAC", "Sidi Kerir/W Med",
              "N. Afr/Euromed", "UK/Cont", "Caribs/USG"]
    pick = {"page_html": html, "html": "", "text": " 280' 725 0 ".join(labels)}
    t = A.picked_table("x", pick)
    assert t is not None and t.title.startswith("Dirty")
    assert A.picked_table("x", {"page_html": html, "html": "", "text": "ничего общего"}) is None


def test_scrub_keeps_field_key():
    from hub.core.claude_pack import scrub_obj
    out = scrub_obj({"fields": [{"key": "price", "name": "x"}], "api_key": "s", "csrf_token": "t"})
    assert out == {"fields": [{"key": "price", "name": "x"}], "api_key": "***", "csrf_token": "***"}
